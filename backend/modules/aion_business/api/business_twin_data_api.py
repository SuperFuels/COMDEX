from __future__ import annotations

from copy import deepcopy
from datetime import UTC, datetime
from hashlib import sha256
import csv
import io
import json
from pathlib import Path
import re
import shutil
from typing import Any, Dict

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field

from backend.modules.aion_business.contracts.business_containers import (
    BusinessContainerMeta,
    BusinessFinancialModelContainer,
    BusinessOperatingModelContainer,
    BusinessIdentityContainer,
    BusinessMapContainer,
    BusinessStructureContainer,
    DepartmentIntelligenceContainer,
)
from backend.modules.aion_business.contracts.workspace import BusinessProfile, WorkspaceSpec
from backend.modules.aion_business.runtime.business_container_repository import BusinessContainerRepository
from backend.modules.aion_business.runtime.canonical_business_identity import canonical_business_aliases, canonical_business_id
from backend.modules.aion_business.runtime.workspace_repository import WorkspaceRepository
from backend.modules.aion_business.runtime.paths import AIONBusinessPaths
from backend.modules.aion_business.runtime.workflow_file_cabinet_repository import WorkflowFileCabinetRepository
from backend.modules.aion_business.runtime.finance_artifact_analyser import analyse_finance_artifact
from backend.modules.aion_business.runtime.finance_completion_service import (
    build_finance_completion,
    seed_operating_model,
)
from backend.modules.aion_business.runtime.finance_director_service import FinanceDirectorService
from backend.modules.aion_business.runtime.business_knowledge_service import BusinessKnowledgeService
from backend.modules.aion_business.runtime.business_map_governance_service import BusinessMapGovernanceService
from backend.modules.aion_business.runtime.business_connector_onboarding_service import BusinessConnectorOnboardingService
from backend.modules.aion_business.runtime.foundation_onboarding_session_service import (
    FoundationOnboardingSessionService,
    FoundationSessionRevisionConflict,
)


router = APIRouter(prefix="/api/aion/business/data", tags=["aion-business-data"])

FINANCE_ARTIFACT_CATEGORIES = {
    "spreadsheets": {".csv", ".xlsx", ".xls", ".xml", ".json"},
    "bank-statements": {".csv", ".xlsx", ".xls", ".pdf", ".ofx", ".qif"},
    "invoices-receipts": {".csv", ".xlsx", ".xls", ".pdf", ".xml", ".json"},
    "management-accounts": {".csv", ".xlsx", ".xls", ".pdf", ".xml", ".json"},
    "vat-tax": {".csv", ".xlsx", ".xls", ".pdf", ".xml", ".json"},
    "payroll": {".csv", ".xlsx", ".xls", ".pdf", ".xml", ".json"},
    "payment-processors": {".csv", ".xlsx", ".xls", ".pdf", ".json"},
    "crm-sales": {".csv", ".xlsx", ".xls", ".pdf", ".json"},
    "reports": {".csv", ".xlsx", ".xls", ".pdf", ".json"},
    "models": {".csv", ".xlsx", ".xls", ".pdf", ".json"},
    "evidence-receipts": {".csv", ".xlsx", ".xls", ".pdf", ".json"},
}
FINANCE_CABINET_FOLDERS = {
    "spreadsheets": "Spreadsheets",
    "bank-statements": "Bank Statements",
    "invoices-receipts": "Invoices & Receipts",
    "management-accounts": "Management Accounts",
    "vat-tax": "VAT & Tax",
    "payroll": "Payroll",
    "payment-processors": "Payment Processors",
    "crm-sales": "CRM & Sales",
    "reports": "Reports",
    "models": "Financial Models",
    "evidence-receipts": "Evidence & Receipts",
}
MAX_FINANCE_ARTIFACT_BYTES = 25 * 1024 * 1024


def _now() -> str:
    return datetime.now(UTC).replace(microsecond=0).isoformat()


def _hash(value: Any) -> str:
    raw = json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return "sha256:" + sha256(raw.encode("utf-8")).hexdigest()


def _object(value: Any) -> Dict[str, Any]:
    return deepcopy(value) if isinstance(value, dict) else {}


def _runtime_department_projection(
    business_id: str,
    departments: Dict[str, Any],
    *,
    business_root: Path | None = None,
) -> Dict[str, Any]:
    """Expose durable Sales and Support runtime evidence to the read-only Board."""
    projected = deepcopy(_object(departments))
    root = business_root or AIONBusinessPaths.business_container_dir(business_id)

    def files(path: Path, pattern: str = "*.json") -> list[Path]:
        return sorted(path.glob(pattern)) if path.exists() else []

    def latest(paths: list[Path]) -> str:
        if not paths:
            return ""
        stamp = max(path.stat().st_mtime for path in paths)
        return datetime.fromtimestamp(stamp, UTC).replace(microsecond=0).isoformat()

    sales_root = root / "sales" / "revenue_spine"
    sales_opportunities = files(sales_root / "opportunities")
    sales_contacts = files(sales_root / "contacts")
    sales_playbooks = files(sales_root / "playbooks")
    sales_simulations = files(sales_root / "simulations")
    sales_calls = files(sales_root / "telephony" / "call_drafts")
    sales_bookings = files(root / "sales" / "completion" / "bookings")
    sales_messages = files(root / "sales" / "completion" / "messages")
    sales_evidence = sales_opportunities + sales_contacts + sales_playbooks + sales_simulations + sales_calls + sales_bookings + sales_messages
    if sales_evidence:
        entry = {**_object(projected.get("sales"))}
        entry.update({
            "department": "sales",
            "status": "operational_foundation_ready",
            "boardroom_summary": (
                f"Canonical Sales runtime contains {len(sales_opportunities)} opportunities, "
                f"{len(sales_contacts)} contacts, {len(sales_playbooks)} playbooks, "
                f"{len(sales_simulations)} simulations, {len(sales_calls)} call drafts, "
                f"{len(sales_bookings)} bookings and {len(sales_messages)} governed messages."
            ),
            "runtime_evidence": {
                "opportunities": len(sales_opportunities), "contacts": len(sales_contacts),
                "playbooks": len(sales_playbooks), "simulations": len(sales_simulations),
                "call_drafts": len(sales_calls), "bookings": len(sales_bookings),
                "messages": len(sales_messages), "external_authority_claimed": False,
            },
            "updated_at": latest(sales_evidence),
        })
        projected["sales"] = entry

    support_root = root / "support" / "case_management"
    support_cases = files(support_root / "cases")
    support_config = [path for path in [support_root / "agent_setup.json", support_root / "polling.json", support_root / "whatsapp.json"] if path.exists()]
    support_voice = files(root / "support" / "voice")
    support_evidence = support_cases + support_config + support_voice
    if support_evidence:
        entry = {**_object(projected.get("support"))}
        entry.update({
            "department": "support",
            "status": "case_authority_foundation_ready",
            "boardroom_summary": (
                f"Canonical Support runtime contains {len(support_cases)} cases, "
                f"{len(support_config)} channel/policy configurations and "
                f"{len(support_voice)} governed voice deployment records."
            ),
            "runtime_evidence": {
                "cases": len(support_cases), "configurations": len(support_config),
                "voice_records": len(support_voice), "automatic_refund_authority_claimed": False,
            },
            "updated_at": latest(support_evidence),
        })
        projected["support"] = entry

    return projected


def _boardroom_operating_projection(model: Dict[str, Any]) -> Dict[str, Any]:
    """Return useful operating facts without exposing confidential HR costing."""
    projected = deepcopy(_object(model))
    projected["labour_resources"] = [
        {key: item.get(key) for key in (
            "id", "person_id", "person_name", "engagement_type", "cost_basis", "available_hours_month",
        )}
        for item in _list(projected.get("labour_resources")) if isinstance(item, dict)
    ]
    projected["job_labour_assignments"] = []
    projected["job_cost_items"] = []
    projected.setdefault("governance", {})["individual_workforce_costs_redacted"] = True
    projected["governance"]["job_economics_are_aggregated"] = True
    return projected


def _list(value: Any) -> list[Any]:
    return deepcopy(value) if isinstance(value, list) else []


def _text(value: Any) -> str | None:
    text = str(value or "").strip()
    return text or None


def _number(value: Any, default: float = 0.0) -> float:
    try:
        if value in (None, ""):
            return default
        return float(str(value).replace(",", "").replace("€", "").replace("£", "").strip())
    except (TypeError, ValueError):
        return default


def _nullable_number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "").replace("€", "").replace("£", "").strip())
    except (TypeError, ValueError):
        return None


def _stable_row_id(prefix: str, row: Dict[str, Any], index: int) -> str:
    existing = str(row.get("id") or "").strip()
    if existing:
        return existing
    seed = str(row.get("sku") or row.get("name") or row.get("description") or index)
    slug = re.sub(r"[^a-z0-9]+", "-", seed.lower()).strip("-") or str(index + 1)
    return f"{prefix}.{slug}"


def _calculate_operating_model(payload: Dict[str, Any]) -> Dict[str, Any]:
    """Normalise editable rows and calculate boardroom-ready operating facts."""
    calculated = deepcopy(payload)
    currency = str(calculated.get("currency") or "EUR").upper()

    inventory: list[Dict[str, Any]] = []
    inventory_value = 0.0
    inventory_available_value = 0.0
    inventory_by_class: Dict[str, float] = {}
    suggested: list[Dict[str, Any]] = []
    for index, raw in enumerate(_list(calculated.get("inventory_items"))):
        item = _object(raw)
        item["id"] = _stable_row_id("inventory", item, index)
        stock_class = str(item.get("stock_class") or "consumable")
        on_hand = _number(item.get("quantity_on_hand"))
        reserved = _number(item.get("quantity_reserved"))
        available = max(on_hand - reserved, 0.0)
        reorder_point = _number(item.get("reorder_point"))
        target = max(_number(item.get("target_stock")), reorder_point)
        unit_cost = _number(item.get("unit_cost"))
        value = on_hand * unit_cost
        item.update({
            "stock_class": stock_class,
            "quantity_on_hand": on_hand,
            "quantity_reserved": reserved,
            "quantity_available": available,
            "unit_cost": unit_cost,
            "stock_value": round(value, 2),
            "reorder_required": available <= reorder_point and target > available,
        })
        inventory.append(item)
        inventory_value += value
        inventory_available_value += available * unit_cost
        inventory_by_class[stock_class] = inventory_by_class.get(stock_class, 0.0) + value
        if item["reorder_required"]:
            qty = max(target - available, 0.0)
            suggested.append({
                "id": f"procurement.auto.{item['id']}",
                "inventory_item_id": item["id"],
                "name": item.get("name") or item.get("sku") or "Stock item",
                "quantity": round(qty, 4),
                "unit": item.get("unit") or "units",
                "supplier_id": item.get("supplier_id"),
                "lead_time_days": _number(item.get("lead_time_days")),
                "estimated_cash_required": round(qty * unit_cost, 2),
                "reason": "available stock is at or below reorder point",
                "status": "suggested",
                "requires_approval": True,
                "generated_by": "stock_policy",
            })
    calculated["inventory_items"] = inventory

    existing_queue = [_object(item) for item in _list(calculated.get("procurement_queue"))]
    existing_ids = {str(item.get("inventory_item_id") or item.get("id") or "") for item in existing_queue}
    existing_queue.extend(item for item in suggested if str(item.get("inventory_item_id")) not in existing_ids)
    for index, item in enumerate(existing_queue):
        item["id"] = _stable_row_id("procurement", item, index)
        item.setdefault("status", "draft")
        item.setdefault("requires_approval", True)
        item["estimated_cash_required"] = round(
            _number(item.get("estimated_cash_required"))
            or _number(item.get("quantity")) * _number(item.get("unit_cost")), 2
        )
    calculated["procurement_queue"] = existing_queue

    bom_by_offering: Dict[str, float] = {}
    for raw in _list(calculated.get("bills_of_materials")):
        line = _object(raw)
        offering_id = str(line.get("offering_id") or "")
        bom_by_offering[offering_id] = bom_by_offering.get(offering_id, 0.0) + (
            _number(line.get("quantity_per_unit"), 1.0) * _number(line.get("unit_cost"))
        )

    economics: list[Dict[str, Any]] = []
    offerings: list[Dict[str, Any]] = []
    for index, raw in enumerate(_list(calculated.get("offerings"))):
        offering = _object(raw)
        offering["id"] = _stable_row_id("offering", offering, index)
        price = _nullable_number(offering.get("price") if offering.get("price") not in (None, "") else offering.get("charge_rate"))
        material = _number(offering.get("material_cost_per_unit")) + bom_by_offering.get(offering["id"], 0.0)
        labour = _number(offering.get("labour_cost_per_unit"))
        other = sum(_number(offering.get(key)) for key in ("equipment_cost_per_unit", "subcontractor_cost_per_unit", "other_direct_cost_per_unit"))
        direct = material + labour + other
        has_direct_cost = bool(bom_by_offering.get(offering["id"])) or any(
            offering.get(key) not in (None, "")
            for key in (
                "material_cost_per_unit", "labour_cost_per_unit", "equipment_cost_per_unit",
                "subcontractor_cost_per_unit", "other_direct_cost_per_unit",
            )
        )
        direct_value = round(direct, 2) if has_direct_cost else None
        contribution = None if price is None or direct_value is None else price - direct_value
        volume = _nullable_number(offering.get("expected_monthly_volume"))
        offering.update({"price": price, "direct_cost_per_unit": direct_value})
        offerings.append(offering)
        economics.append({
            "offering_id": offering["id"],
            "name": offering.get("name") or offering.get("sku") or "Offering",
            "currency": currency,
            "price_per_unit": None if price is None else round(price, 2),
            "direct_cost_per_unit": direct_value,
            "contribution_per_unit": None if contribution is None else round(contribution, 2),
            "gross_margin_percent": round((contribution / price * 100), 2) if contribution is not None and price else None,
            "expected_monthly_volume": volume,
            "expected_monthly_revenue": None if price is None or volume is None else round(price * volume, 2),
            "expected_monthly_contribution": None if contribution is None or volume is None else round(contribution * volume, 2),
        })
    calculated["offerings"] = offerings
    calculated["unit_economics"] = economics

    labour_resources: list[Dict[str, Any]] = []
    labour_by_id: Dict[str, Dict[str, Any]] = {}
    labour_by_person: Dict[str, Dict[str, Any]] = {}
    for index, raw in enumerate(_list(calculated.get("labour_resources"))):
        resource = _object(raw)
        resource["id"] = _stable_row_id("labour", resource, index)
        resource["cost_rate"] = _nullable_number(resource.get("cost_rate"))
        labour_resources.append(resource)
        labour_by_id[resource["id"]] = resource
        if resource.get("person_id"):
            labour_by_person[str(resource["person_id"])] = resource
    calculated["labour_resources"] = labour_resources

    jobs: list[Dict[str, Any]] = []
    job_ids: set[str] = set()
    for index, raw in enumerate(_list(calculated.get("jobs"))):
        job = _object(raw)
        job["id"] = _stable_row_id("job", job, index)
        job["status"] = str(job.get("status") or "estimate")
        job["estimated_revenue"] = _nullable_number(job.get("estimated_revenue"))
        job["actual_revenue"] = _nullable_number(job.get("actual_revenue"))
        jobs.append(job)
        job_ids.add(job["id"])
    calculated["jobs"] = jobs

    assignments: list[Dict[str, Any]] = []
    for index, raw in enumerate(_list(calculated.get("job_labour_assignments"))):
        assignment = _object(raw)
        if str(assignment.get("job_id") or "") not in job_ids:
            continue
        assignment["id"] = _stable_row_id("job-labour", assignment, index)
        resource = labour_by_id.get(str(assignment.get("labour_resource_id") or "")) or labour_by_person.get(str(assignment.get("person_id") or "")) or {}
        basis = str(assignment.get("cost_basis") or resource.get("cost_basis") or "hourly")
        rate = _number(assignment.get("cost_rate") if assignment.get("cost_rate") not in (None, "") else resource.get("cost_rate"))
        estimated_units = _number(assignment.get("estimated_units"))
        actual_units = _number(assignment.get("actual_units"))
        assignment.update({
            "person_id": assignment.get("person_id") or resource.get("person_id"),
            "person_name": assignment.get("person_name") or resource.get("person_name") or resource.get("name"),
            "cost_basis": basis, "cost_rate": round(rate, 2),
            "estimated_units": estimated_units, "actual_units": actual_units,
            "estimated_cost": round(rate * estimated_units, 2),
            "actual_cost": round(rate * actual_units, 2),
        })
        assignments.append(assignment)
    calculated["job_labour_assignments"] = assignments

    cost_items: list[Dict[str, Any]] = []
    for index, raw in enumerate(_list(calculated.get("job_cost_items"))):
        item = _object(raw)
        if str(item.get("job_id") or "") not in job_ids:
            continue
        item["id"] = _stable_row_id("job-cost", item, index)
        item["category"] = str(item.get("category") or "other_direct")
        item["estimated_cost"] = round(_number(item.get("estimated_cost")), 2)
        item["actual_cost"] = round(_number(item.get("actual_cost")), 2)
        cost_items.append(item)
    calculated["job_cost_items"] = cost_items

    job_economics: list[Dict[str, Any]] = []
    for job in jobs:
        job_id = job["id"]
        labour_estimate = sum(_number(item.get("estimated_cost")) for item in assignments if item.get("job_id") == job_id)
        labour_actual = sum(_number(item.get("actual_cost")) for item in assignments if item.get("job_id") == job_id)
        other_estimate = sum(_number(item.get("estimated_cost")) for item in cost_items if item.get("job_id") == job_id)
        other_actual = sum(_number(item.get("actual_cost")) for item in cost_items if item.get("job_id") == job_id)
        estimated_revenue = job.get("estimated_revenue")
        actual_revenue = job.get("actual_revenue")
        estimated_cost = labour_estimate + other_estimate
        actual_cost = labour_actual + other_actual
        estimated_contribution = None if estimated_revenue is None else estimated_revenue - estimated_cost
        actual_contribution = None if actual_revenue is None else actual_revenue - actual_cost
        job_economics.append({
            "job_id": job_id, "name": job.get("name") or "Job / quote", "status": job.get("status"),
            "estimated_revenue": estimated_revenue, "estimated_labour_cost": round(labour_estimate, 2),
            "estimated_other_direct_cost": round(other_estimate, 2), "estimated_direct_cost": round(estimated_cost, 2),
            "estimated_contribution": None if estimated_contribution is None else round(estimated_contribution, 2),
            "estimated_margin_percent": round(estimated_contribution / estimated_revenue * 100, 2) if estimated_revenue and estimated_contribution is not None else None,
            "actual_revenue": actual_revenue, "actual_labour_cost": round(labour_actual, 2),
            "actual_other_direct_cost": round(other_actual, 2), "actual_direct_cost": round(actual_cost, 2),
            "actual_contribution": None if actual_contribution is None else round(actual_contribution, 2),
            "actual_margin_percent": round(actual_contribution / actual_revenue * 100, 2) if actual_revenue and actual_contribution is not None else None,
            "cost_variance": round(actual_cost - estimated_cost, 2),
        })
    calculated["job_economics"] = job_economics
    calculated["inventory_metrics"] = {
        "currency": currency,
        "total_stock_value": round(inventory_value, 2),
        "available_stock_value": round(inventory_available_value, 2),
        "stock_value_by_class": {key: round(value, 2) for key, value in inventory_by_class.items()},
        "reorder_items": len(suggested),
        "suggested_procurement_cash_required": round(sum(_number(item.get("estimated_cash_required")) for item in suggested), 2),
    }
    calculated["currency"] = currency
    calculated["model_status"] = (
        "needs_user_review"
        if str(payload.get("model_status") or "") == "needs_user_review"
        else "ready" if offerings else "draft"
    )
    return calculated


def _project_operating_model_to_business_map(repository: BusinessContainerRepository, business_id: str, model: Dict[str, Any], source: str) -> None:
    try:
        business_map = repository.load_business_map(business_id)
    except FileNotFoundError:
        business_map = BusinessMapContainer(
            id=f"{business_id}.business_map", workspace_id=business_id,
            meta=_meta(business_id, "business_map", source),
        )
    business_map.facts = [
        fact for fact in business_map.facts
        if not str(_object(fact).get("source_ref") or "").startswith("operating_model:")
    ]
    projected = {
        "offerings": _list(model.get("offerings")),
        "unit_economics": _list(model.get("unit_economics")),
        "inventory_metrics": _object(model.get("inventory_metrics")),
        "procurement_queue": _list(model.get("procurement_queue")),
        "production_lines": _list(model.get("production_lines")),
        "labour_resources": [
            {key: item.get(key) for key in ("id", "person_id", "person_name", "engagement_type", "cost_basis", "available_hours_month")}
            for item in _list(model.get("labour_resources")) if isinstance(item, dict)
        ],
        "job_economics": _list(model.get("job_economics")),
        "payment_terms": _list(model.get("payment_terms")),
        "financial_targets": _list(model.get("financial_targets")),
    }
    observed_at = _now()
    for field, value in projected.items():
        business_map.facts.append({
            "business_id": business_id, "function": "operating_model", "category": "business_model",
            "field": field, "value": value, "classification": "structured_business_fact",
            "confidence": 0.94, "verification_status": "user_confirmed",
            "source_ref": f"operating_model:{field}", "observed_at": observed_at, "created_at": observed_at,
        })
    business_map.revision = int(business_map.revision or 0) + 1
    business_map.meta = _meta(business_id, "business_map", source)
    repository.save_model(business_map)


def _business_type(value: Any) -> str:
    raw = str(value or "").strip().lower().replace(" ", "_")
    aliases = {
        "service": "service_business",
        "services": "service_business",
        "service_business": "service_business",
        "ecommerce": "ecommerce",
        "e-commerce": "ecommerce",
        "marketplace": "marketplace_seller",
        "marketplace_seller": "marketplace_seller",
        "agency": "agency",
        "hybrid": "hybrid",
        "mixed": "hybrid",
        "mixture": "hybrid",
    }
    return aliases.get(raw, "service_business")


class FoundationCommitRequest(BaseModel):
    packet: Dict[str, Any]
    source: str = "business_foundation"


class FoundationSessionWriteRequest(BaseModel):
    packet: Dict[str, Any]
    expected_revision: int | None = None
    actor_id: str = Field(default="local_owner", min_length=1, max_length=160)
    device_id: str = Field(default="unknown_device", min_length=1, max_length=160)
    complete: bool = False


class FoundationSessionResetRequest(BaseModel):
    expected_revision: int | None = None
    actor_id: str = Field(default="local_owner", min_length=1, max_length=160)
    confirm_reset: bool = False


class ContainerWriteRequest(BaseModel):
    payload: Dict[str, Any]
    source: str = "desktop_runtime"
    expected_revision: int | None = None


class BusinessMapFactsRequest(BaseModel):
    facts: list[Dict[str, Any]] = Field(default_factory=list)
    source: str = "desktop_runtime"


class BusinessMapFactProposalRequest(BaseModel):
    field: str = Field(..., min_length=1, max_length=160)
    value: Any
    source_ref: str = Field(..., min_length=1, max_length=500)
    confidence: float = Field(..., ge=0, le=1)
    proposed_by: str = Field(..., min_length=1, max_length=160)
    function: str = Field(default="business", max_length=100)
    category: str = Field(default="foundation", max_length=100)
    review_owner_id: str | None = Field(default=None, max_length=160)
    effective_from: str | None = None
    effective_until: str | None = None
    expected_revision: int | None = None


class BusinessMapFactReviewRequest(BaseModel):
    action: str
    actor_id: str = Field(..., min_length=1, max_length=160)
    expected_revision: int | None = None
    corrected_value: Any = None
    reason: str = Field(default="", max_length=500)
    confirm_delete: bool = False


class BusinessMapFactMergeRequest(BaseModel):
    primary_fact_id: str = Field(..., min_length=1, max_length=160)
    duplicate_fact_id: str = Field(..., min_length=1, max_length=160)
    actor_id: str = Field(..., min_length=1, max_length=160)
    expected_revision: int | None = None


class BusinessMapRelationshipProposalRequest(BaseModel):
    from_id: str = Field(..., min_length=1, max_length=240)
    relationship_type: str = Field(..., min_length=1, max_length=160)
    to_id: str = Field(..., min_length=1, max_length=240)
    source_ref: str = Field(..., min_length=1, max_length=500)
    confidence: float = Field(..., ge=0, le=1)
    proposed_by: str = Field(..., min_length=1, max_length=160)
    review_owner_id: str | None = Field(default=None, max_length=160)
    expected_revision: int | None = None


class BusinessMapRelationshipReviewRequest(BaseModel):
    action: str
    actor_id: str = Field(..., min_length=1, max_length=160)
    expected_revision: int | None = None
    corrected_from_id: str | None = Field(default=None, max_length=240)
    corrected_relationship_type: str | None = Field(default=None, max_length=160)
    corrected_to_id: str | None = Field(default=None, max_length=240)
    reason: str = Field(default="", max_length=500)
    confirm_delete: bool = False


class BusinessMapRelationshipNormalizeRequest(BaseModel):
    actor_id: str = Field(..., min_length=1, max_length=160)
    expected_revision: int | None = None


class FinanceAgentTurnRequest(BaseModel):
    user_text: str = Field(..., min_length=1, max_length=4000)


class FinanceArtifactAcceptanceRequest(BaseModel):
    fact_ids: list[str] = Field(default_factory=list)


class BusinessKnowledgeTextRequest(BaseModel):
    title: str = Field(default="Owner-provided business knowledge", max_length=240)
    text: str = Field(..., min_length=1, max_length=500_000)
    scopes: list[str] = Field(default_factory=lambda: ["all"])
    confidentiality: str = "internal"
    effective_from: str | None = None
    effective_until: str | None = None


class BusinessKnowledgeApprovalRequest(BaseModel):
    claim_ids: list[str] = Field(default_factory=list)
    actor: str = "business_owner"


class BusinessKnowledgeSearchRequest(BaseModel):
    query: str = Field(..., min_length=1, max_length=4000)
    actor_scope: str = "pilot"
    limit: int = Field(default=8, ge=1, le=25)


class BusinessKnowledgeEntityLinkRequest(BaseModel):
    actor: str = Field(default="business_owner", min_length=1, max_length=160)


class BusinessKnowledgeSourcePolicyRequest(BaseModel):
    actor: str = Field(default="business_owner", min_length=1, max_length=160)
    legal_hold: bool | None = None
    retention_until: str | None = None


class BusinessKnowledgeSourceRevokeRequest(BaseModel):
    actor: str = Field(default="business_owner", min_length=1, max_length=160)
    confirm_revoke: bool = False
    reason: str = Field(default="", max_length=500)


class BusinessKnowledgeRetentionRequest(BaseModel):
    actor: str = Field(default="business_owner", min_length=1, max_length=160)


def _repo() -> BusinessContainerRepository:
    return BusinessContainerRepository()


def _workspaces() -> WorkspaceRepository:
    return WorkspaceRepository()


def _business_knowledge() -> BusinessKnowledgeService:
    return BusinessKnowledgeService()


def _business_map_governance() -> BusinessMapGovernanceService:
    return BusinessMapGovernanceService(_repo())


def _connector_onboarding() -> BusinessConnectorOnboardingService:
    return BusinessConnectorOnboardingService()


def _ensure_workspace(business_id: str, draft: Dict[str, Any]) -> None:
    repository = _workspaces()
    try:
        repository.load(business_id)
        return
    except FileNotFoundError:
        pass

    business_name = _text(draft.get("business_name") or draft.get("name")) or business_id
    repository.save(
        WorkspaceSpec(
            id=business_id,
            name=business_name,
            business_type=_business_type(draft.get("business_type") or draft.get("business_model")),
            owner=_text(draft.get("user_name") or draft.get("owner") or draft.get("role")) or "owner",
            deployment_mode="local",
            business_profile=BusinessProfile(
                legal_name=business_name,
                trading_name=business_name,
                sector=_text(draft.get("industry") or draft.get("business_type")),
                stage=_text(draft.get("business_stage") or draft.get("stage")),
                products=[_text(draft.get("offer_summary") or draft.get("products_services"))]
                if _text(draft.get("offer_summary") or draft.get("products_services")) else [],
                customer_types=[_text(draft.get("target_customers"))] if _text(draft.get("target_customers")) else [],
                preferred_tools=[_text(draft.get("current_tools") or draft.get("systems"))]
                if _text(draft.get("current_tools") or draft.get("systems")) else [],
            ),
            status="active",
        )
    )


def _migrate_legacy_aliases(candidate: str, business_id: str) -> list[Dict[str, Any]]:
    """Copy legacy underscore-scoped runtime data into the canonical hyphen scope.

    Legacy data is retained. This is deliberately non-destructive until a later
    verified cleanup migration is approved.
    """
    migrations: list[Dict[str, Any]] = []
    for alias in canonical_business_aliases(candidate):
        if alias == business_id:
            continue
        legacy_container = AIONBusinessPaths.BUSINESS_CONTAINERS / alias
        canonical_container = AIONBusinessPaths.BUSINESS_CONTAINERS / business_id
        if legacy_container.exists():
            canonical_container.mkdir(parents=True, exist_ok=True)
            shutil.copytree(legacy_container, canonical_container, dirs_exist_ok=True)
            for copied_json in canonical_container.glob("*.json"):
                try:
                    payload = json.loads(copied_json.read_text(encoding="utf-8"))
                    if isinstance(payload, dict):
                        payload["workspace_id"] = business_id
                        if str(payload.get("id") or "").startswith(f"{alias}."):
                            payload["id"] = f"{business_id}.{str(payload['id']).split('.', 1)[1]}"
                        if isinstance(payload.get("meta"), dict):
                            payload["meta"]["workspace_id"] = business_id
                            payload["meta"]["source"] = "legacy_alias_migration"
                            payload["meta"]["updated_at"] = _now()
                        copied_json.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")
                except Exception:
                    continue
            migrations.append({"kind": "business_container", "from": alias, "to": business_id})

        legacy_cabinet = AIONBusinessPaths.ROOT / "workflow_file_cabinets" / alias / "tree.json"
        canonical_cabinet = AIONBusinessPaths.ROOT / "workflow_file_cabinets" / business_id / "tree.json"
        if legacy_cabinet.exists() and not canonical_cabinet.exists():
            canonical_cabinet.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(legacy_cabinet, canonical_cabinet)
            try:
                tree = json.loads(canonical_cabinet.read_text(encoding="utf-8"))
                tree["workspace_id"] = business_id
                tree["business_container"] = business_id
                canonical_cabinet.write_text(json.dumps(tree, indent=2, sort_keys=True), encoding="utf-8")
            except Exception:
                pass
            migrations.append({"kind": "file_cabinet", "from": alias, "to": business_id})

        legacy_workspace = AIONBusinessPaths.WORKSPACES / f"{alias}.json"
        canonical_workspace = AIONBusinessPaths.WORKSPACES / f"{business_id}.json"
        if legacy_workspace.exists() and not canonical_workspace.exists():
            canonical_workspace.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(legacy_workspace, canonical_workspace)
            migrations.append({"kind": "workspace", "from": alias, "to": business_id})
    return migrations


def _meta(business_id: str, key: str, source: str) -> BusinessContainerMeta:
    return BusinessContainerMeta(
        workspace_id=business_id,
        container_key=key,
        updated_at=_now(),
        source=source,
    )


def _receipt(*, business_id: str, kind: str, payload: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "business_id": business_id,
        "container_kind": kind,
        "revision": int(payload.get("revision") or 1),
        "payload_hash": _hash(payload),
        "written_at": _now(),
        "storage_state": "committed_to_business_container",
    }


def _safe_filename(value: str) -> str:
    name = Path(str(value or "upload")).name.strip()
    stem = re.sub(r"[^a-zA-Z0-9_. -]+", "-", name).strip(". ")
    return stem[:180] or "upload"


def _ensure_folder(children: list[Dict[str, Any]], *, node_id: str, name: str, node_type: str = "folder") -> Dict[str, Any]:
    node = next((item for item in children if isinstance(item, dict) and item.get("id") == node_id), None)
    if node is None:
        node = {"id": node_id, "name": name, "type": node_type, "children": []}
        children.append(node)
    node["name"] = name
    node["type"] = node_type
    if not isinstance(node.get("children"), list):
        node["children"] = []
    return node


def _add_finance_file_cabinet_pointer(*, business_id: str, category: str, record: Dict[str, Any]) -> Dict[str, Any]:
    tree = WorkflowFileCabinetRepository.load(business_id)
    folders = tree.get("folders")
    if not isinstance(folders, list):
        folders = list((tree.get("root") or {}).get("children") or [])
        tree["folders"] = folders

    finance = _ensure_folder(folders, node_id="folder_finance", name="Finance", node_type="department")
    for key, label in FINANCE_CABINET_FOLDERS.items():
        _ensure_folder(finance["children"], node_id=f"folder_finance_{key.replace('-', '_')}", name=label)
    destination = _ensure_folder(
        finance["children"],
        node_id=f"folder_finance_{category.replace('-', '_')}",
        name=FINANCE_CABINET_FOLDERS[category],
    )

    pointer_id = f"finance_artifact_pointer_{record['artifact_id']}"
    pointer = {
        "id": pointer_id,
        "name": record["target"]["artifact_name"],
        "type": "business_container_artifact",
        "document_type": record["target"]["artifact_type"],
        "source_of_truth": "business_container",
        "file_cabinet_role": "index_pointer_only",
        "status": "committed",
        "provenance": {"created_by": "user_upload"},
        "deletion_policy": "user_removable",
        "target": {
            "business_container_id": business_id,
            "sub_container": "finance",
            "category": category,
            "artifact_type": record["target"]["artifact_type"],
            "storage_path": record["storage_path"],
            "artifact_id": record["artifact_id"],
            "artifact_hash": record["artifact_hash"],
            "receipt_hash": record["receipt_hash"],
            "replay_locator_hash": record["replay_locator_hash"],
            "analysis_status": _object(record.get("analysis")).get("status"),
            "analysis_path": _object(record.get("analysis")).get("storage_path"),
            "analysis_hash": _object(record.get("analysis")).get("analysis_hash"),
        },
    }
    existing_index = next((index for index, item in enumerate(destination["children"]) if item.get("id") == pointer_id), None)
    if existing_index is None:
        destination["children"].append(pointer)
    else:
        destination["children"][existing_index] = pointer
    tree["root"] = {**_object(tree.get("root")), "id": "root", "type": "folder", "name": "Workflows", "children": folders}
    WorkflowFileCabinetRepository.save(business_id, tree)
    return pointer


def _add_operating_model_file_pointer(*, business_id: str, category: str, record: Dict[str, Any]) -> Dict[str, Any]:
    tree = WorkflowFileCabinetRepository.load(business_id)
    folders = tree.get("folders")
    if not isinstance(folders, list):
        folders = list((tree.get("root") or {}).get("children") or [])
        tree["folders"] = folders
    department = _ensure_folder(folders, node_id="folder_products_services", name="Products & Services", node_type="department")
    destination = _ensure_folder(
        department["children"], node_id=f"folder_products_services_{category}",
        name="Offer Catalogue" if category == "offerings" else "Stock & Inventory",
    )
    pointer = {
        "id": f"operating_model_import_{record['artifact_id']}",
        "name": record["name"], "type": "business_container_artifact",
        "document_type": record["extension"], "source_of_truth": "business_container",
        "file_cabinet_role": "index_pointer_only", "status": "committed",
        "provenance": {"created_by": "user_upload"}, "deletion_policy": "user_removable",
        "target": {
            "business_container_id": business_id, "sub_container": "operating_model",
            "category": category, "storage_path": record["storage_path"],
            "artifact_id": record["artifact_id"], "artifact_hash": record["artifact_hash"],
            "imported_row_count": record["imported_row_count"],
        },
    }
    destination["children"] = [item for item in destination["children"] if item.get("id") != pointer["id"]] + [pointer]
    tree["root"] = {**_object(tree.get("root")), "id": "root", "type": "folder", "name": "Workflows", "children": folders}
    WorkflowFileCabinetRepository.save(business_id, tree)
    return pointer


def _normalise_import_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def _parse_operating_model_rows(content: bytes, extension: str) -> list[Dict[str, Any]]:
    if extension == ".csv":
        decoded = content.decode("utf-8-sig", errors="replace")
        return [{_normalise_import_header(key): value for key, value in row.items()} for row in csv.DictReader(io.StringIO(decoded))]
    if extension == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        headers = [_normalise_import_header(value) for value in next(iterator, [])]
        return [{headers[index]: value for index, value in enumerate(row) if index < len(headers) and headers[index]} for row in iterator]
    raise HTTPException(status_code=400, detail="operating_model_import_requires_csv_or_xlsx")


def _first_value(row: Dict[str, Any], *names: str) -> Any:
    for name in names:
        if row.get(name) not in (None, ""):
            return row[name]
    return None


def _map_operating_import_rows(rows: list[Dict[str, Any]], category: str) -> list[Dict[str, Any]]:
    mapped: list[Dict[str, Any]] = []
    for row in rows:
        if category == "offerings":
            name = _first_value(row, "name", "offering", "product", "service", "product_name", "service_name", "description")
            if not name:
                continue
            mapped.append({
                "name": str(name), "sku": _first_value(row, "sku", "product_code", "code"),
                "offering_type": str(_first_value(row, "type", "offering_type", "product_type") or "product"),
                "pricing_basis": str(_first_value(row, "pricing_basis", "charge_basis", "rate_type") or "per_unit"),
                "price": _number(_first_value(row, "price", "rate", "charge_rate", "sell_price", "unit_price")),
                "unit": str(_first_value(row, "unit", "charge_unit", "uom") or "units"),
                "expected_monthly_volume": _number(_first_value(row, "expected_monthly_volume", "monthly_volume", "volume")),
                "material_cost_per_unit": _number(_first_value(row, "material_cost_per_unit", "material_cost", "unit_cost", "cost")),
                "labour_cost_per_unit": _number(_first_value(row, "labour_cost_per_unit", "labour_cost")),
                "import_source": "catalogue_spreadsheet",
            })
        else:
            name = _first_value(row, "name", "item", "stock_item", "product", "description")
            if not name:
                continue
            mapped.append({
                "name": str(name), "sku": _first_value(row, "sku", "product_code", "code"),
                "stock_class": str(_first_value(row, "stock_class", "inventory_type", "type") or "raw_material"),
                "unit": str(_first_value(row, "unit", "uom") or "units"),
                "quantity_on_hand": _number(_first_value(row, "quantity_on_hand", "on_hand", "quantity", "stock")),
                "quantity_reserved": _number(_first_value(row, "quantity_reserved", "reserved")),
                "unit_cost": _number(_first_value(row, "unit_cost", "cost", "cost_price")),
                "sell_price": _number(_first_value(row, "sell_price", "price", "unit_price")),
                "reorder_point": _number(_first_value(row, "reorder_point", "minimum_stock", "min_stock")),
                "target_stock": _number(_first_value(row, "target_stock", "maximum_stock", "max_stock")),
                "supplier_name": _first_value(row, "supplier_name", "supplier"),
                "lead_time_days": _number(_first_value(row, "lead_time_days", "lead_time")),
                "import_source": "inventory_spreadsheet",
            })
    return mapped


@router.get("/identity/{candidate}")
def resolve_business_identity(candidate: str) -> Dict[str, Any]:
    try:
        business_id = canonical_business_id(candidate)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {
        "ok": True,
        "business_id": business_id,
        "workspace_id": business_id,
        "container_id": business_id,
        "aliases": canonical_business_aliases(candidate),
    }


def _foundation_sessions() -> FoundationOnboardingSessionService:
    return FoundationOnboardingSessionService()


@router.get("/foundation-session/{candidate}")
def get_foundation_session(candidate: str) -> Dict[str, Any]:
    try:
        session = _foundation_sessions().get(candidate)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True, "session": session}


@router.put("/foundation-session/{candidate}")
def put_foundation_session(
    candidate: str, request: FoundationSessionWriteRequest
) -> Dict[str, Any]:
    try:
        session = _foundation_sessions().save(
            candidate,
            packet=request.packet,
            expected_revision=request.expected_revision,
            actor_id=request.actor_id,
            device_id=request.device_id,
            complete=request.complete,
        )
    except FoundationSessionRevisionConflict as exc:
        raise HTTPException(
            status_code=409,
            detail={
                "reason": "foundation_session_revision_conflict",
                "expected": exc.expected,
                "current": exc.current,
            },
        ) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True, "session": session}


@router.post("/foundation-session/{candidate}/reset")
def reset_foundation_session(
    candidate: str, request: FoundationSessionResetRequest
) -> Dict[str, Any]:
    try:
        session = _foundation_sessions().reset(
            candidate,
            expected_revision=request.expected_revision,
            actor_id=request.actor_id,
            confirm_reset=request.confirm_reset,
        )
    except FoundationSessionRevisionConflict as exc:
        raise HTTPException(
            status_code=409,
            detail={
                "reason": "foundation_session_revision_conflict",
                "expected": exc.expected,
                "current": exc.current,
            },
        ) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True, "session": session}


@router.post("/foundation/{candidate}")
def commit_foundation(candidate: str, request: FoundationCommitRequest) -> Dict[str, Any]:
    try:
        business_id = canonical_business_id(candidate)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    packet = _object(request.packet)
    draft = _object(packet.get("foundation_draft"))
    business_map = _object(packet.get("business_map"))
    migrations = _migrate_legacy_aliases(candidate, business_id)
    _ensure_workspace(business_id, draft)
    repository = _repo()

    name = _text(draft.get("business_name") or draft.get("name")) or business_id
    identity = BusinessIdentityContainer(
        id=f"{business_id}.business_identity",
        workspace_id=business_id,
        meta=_meta(business_id, "business_identity", request.source),
        legal_name=name,
        trading_name=name,
        business_type=_business_type(draft.get("business_type") or draft.get("business_model")),
        sector=_text(draft.get("industry") or draft.get("business_type")),
        stage=_text(draft.get("business_stage") or draft.get("stage")),
        owner=_text(draft.get("user_name") or draft.get("owner") or draft.get("role")),
        country=_text(draft.get("country")) or "unspecified",
        region=_text(draft.get("region") or draft.get("service_area") or draft.get("location")),
        city=_text(draft.get("city")),
        timezone=_text(draft.get("timezone")) or "UTC",
        currency=_text(draft.get("currency")) or "unspecified",
        primary_domain=_text(draft.get("primary_domain")),
        website_url=_text(draft.get("website_url") or draft.get("website")),
        primary_email=_text(draft.get("contact_email") or draft.get("primary_email")),
        description=_text(draft.get("offer_summary") or draft.get("products_services")),
        source_refs=[request.source, str(packet.get("schema_version") or "foundation_packet")],
    )
    repository.save_model(identity)

    structure = BusinessStructureContainer(
        id=f"{business_id}.business_structure",
        workspace_id=business_id,
        meta=_meta(business_id, "business_structure", request.source),
        identity={
            "workspace_id": business_id,
            "business_name": name,
            "legal_name": name,
            "public_brand_name": name,
            "business_type": _business_type(draft.get("business_type") or draft.get("business_model")),
            "country": _text(draft.get("country")) or "unspecified",
            "region": _text(draft.get("region") or draft.get("service_area") or draft.get("location")),
            "city": _text(draft.get("city")),
            "timezone": _text(draft.get("timezone")) or "UTC",
            "currency": _text(draft.get("currency")) or "unspecified",
            "primary_domain": _text(draft.get("primary_domain")),
            "website_url": _text(draft.get("website_url") or draft.get("website")),
            "primary_email": _text(draft.get("contact_email") or draft.get("primary_email")),
            "description": _text(draft.get("offer_summary") or draft.get("products_services")),
            "founder_name": _text(draft.get("user_name") or draft.get("owner")),
            "onboarding_status": "foundation_complete",
            "active": True,
        },
        services=[{
            "id": f"{business_id}.service.foundation",
            "workspace_id": business_id,
            "name": _text(draft.get("offer_summary") or draft.get("products_services")) or "Primary service",
            "slug": "primary-service",
            "description": _text(draft.get("offer_summary") or draft.get("products_services")),
            "delivery_mode": "onsite",
            "pricing_type": "quote",
            "currency": _text(draft.get("currency")) or "EUR",
            "service_area": [_text(draft.get("service_area") or draft.get("location"))] if _text(draft.get("service_area") or draft.get("location")) else [],
            "status": "active",
            "active": True,
        }] if draft.get("offer_summary") or draft.get("products_services") else [],
        revenue_streams=[{
            "id": f"{business_id}.revenue.foundation",
            "workspace_id": business_id,
            "name": "Primary revenue stream",
            "stream_type": "service_sales",
            "cadence": "variable",
            "status": "active",
            "description": _text(draft.get("revenue_model") or draft.get("business_model")),
            "currency": _text(draft.get("currency")) or "EUR",
            "active": True,
        }] if draft.get("revenue_model") or draft.get("business_model") else [],
        fulfillment_processes=[{
            "id": f"{business_id}.fulfillment.foundation",
            "workspace_id": business_id,
            "name": "Primary delivery process",
            "fulfillment_mode": "onsite_service",
            "status": "active",
            "description": _text(draft.get("delivery_model")),
            "active": True,
        }] if draft.get("delivery_model") else [],
        functions=[{
            "id": f"{business_id}.function.{key}",
            "workspace_id": business_id,
            "department_key": key,
            "key": key,
            "name": key.title(),
            "description": f"{key.title()} business function",
            "status": "draft",
            "active": True,
        } for key in ["products_services", "finance", "marketing", "sales", "operations", "support", "hr", "procurement"]],
        teams=[],
    )
    repository.save_model(structure)

    map_model = BusinessMapContainer(
        id=f"{business_id}.business_map",
        workspace_id=business_id,
        meta=_meta(business_id, "business_map", request.source),
        facts=_list(business_map.get("facts")),
        assumptions=_list(business_map.get("assumptions")),
        unknowns=_list(business_map.get("unknowns")),
        conflicts=_list(business_map.get("conflicts")),
        unanswered_fields=[str(x) for x in _list(business_map.get("unanswered_fields"))],
        department_discovery_gaps=[str(x) for x in _list(business_map.get("department_discovery_gaps"))],
        source_refs=[request.source, str(packet.get("schema_version") or "foundation_packet")],
        revision=1,
    )
    repository.save_model(map_model)

    try:
        ledger = repository.load_department_intelligence(business_id)
    except FileNotFoundError:
        ledger = DepartmentIntelligenceContainer(
            id=f"{business_id}.department_intelligence",
            workspace_id=business_id,
            meta=_meta(business_id, "department_intelligence", request.source),
            departments={key: {"department": key, "status": "needs_discovery"} for key in ["products_services", "finance", "marketing", "sales", "operations", "support", "hr", "procurement"]},
        )
        repository.save_model(ledger)

    return {
        "ok": True,
        "business_id": business_id,
        "legacy_alias_migrations": migrations,
        "containers": {
            "business_identity": _receipt(business_id=business_id, kind="business_identity", payload=identity.model_dump(mode="json")),
            "business_structure": _receipt(business_id=business_id, kind="business_structure", payload=structure.model_dump(mode="json")),
            "business_map": _receipt(business_id=business_id, kind="business_map", payload=map_model.model_dump(mode="json")),
            "department_intelligence": _receipt(business_id=business_id, kind="department_intelligence", payload=ledger.model_dump(mode="json")),
        },
    }


def _read_container(candidate: str, kind: str) -> Dict[str, Any]:
    try:
        business_id = canonical_business_id(candidate)
        payload = _repo().load_dict(business_id, kind)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail=f"{kind}_not_found") from exc
    return {"ok": True, "business_id": business_id, "payload": payload, "receipt": _receipt(business_id=business_id, kind=kind, payload=payload)}


def _next_revision(existing: Dict[str, Any] | None, expected: int | None) -> int:
    current = int((existing or {}).get("revision") or 0)
    if expected is not None and current != expected:
        raise HTTPException(status_code=409, detail={"reason": "revision_conflict", "expected": expected, "current": current})
    return current + 1


@router.get("/business-map/{candidate}")
def get_business_map(candidate: str) -> Dict[str, Any]:
    return _read_container(candidate, "business_map")


@router.get("/connector-onboarding/{candidate}")
def get_business_connector_onboarding(candidate: str) -> Dict[str, Any]:
    try:
        return _connector_onboarding().snapshot(candidate)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@router.put("/business-map/{candidate}")
def put_business_map(candidate: str, request: ContainerWriteRequest) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    repository = _repo()
    existing = repository.load_optional_dict(business_id, "business_map")
    payload = _object(request.payload)
    payload.update({"id": f"{business_id}.business_map", "workspace_id": business_id, "kind": "business_map"})
    payload["revision"] = _next_revision(existing, request.expected_revision)
    payload["meta"] = _meta(business_id, "business_map", request.source).model_dump(mode="json")
    model = BusinessMapContainer(**payload)
    repository.save_model(model)
    data = model.model_dump(mode="json")
    return {"ok": True, "business_id": business_id, "payload": data, "receipt": _receipt(business_id=business_id, kind="business_map", payload=data)}


@router.post("/business-map/{candidate}/facts")
def append_business_map_facts(candidate: str, request: BusinessMapFactsRequest) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    repository = _repo()
    try:
        model = repository.load_business_map(business_id)
    except FileNotFoundError:
        model = BusinessMapContainer(
            id=f"{business_id}.business_map",
            workspace_id=business_id,
            meta=_meta(business_id, "business_map", request.source),
        )

    existing_keys = {
        (str(item.get("function") or ""), str(item.get("field") or ""), json.dumps(item.get("value"), sort_keys=True, default=str), str(item.get("source_ref") or ""))
        for item in model.facts if isinstance(item, dict)
    }
    appended = 0
    for incoming in request.facts:
        fact = _object(incoming)
        fact.setdefault("business_id", business_id)
        fact.setdefault("classification", "founder_provided_fact")
        fact.setdefault("confidence", 0.88)
        fact.setdefault("verification_status", "unverified")
        fact.setdefault("observed_at", _now())
        fact.setdefault("created_at", _now())
        key = (str(fact.get("function") or ""), str(fact.get("field") or ""), json.dumps(fact.get("value"), sort_keys=True, default=str), str(fact.get("source_ref") or ""))
        if key in existing_keys:
            continue
        model.facts.append(fact)
        existing_keys.add(key)
        appended += 1

    # A no-op sync must not rewrite metadata and silently change the payload
    # hash while leaving the semantic revision unchanged. Boardroom approvals
    # bind to both revision and hash, so only a real append may mutate the map.
    if appended:
        model.revision = int(model.revision or 0) + 1
        model.meta.updated_at = _now()
        model.meta.source = request.source
        repository.save_model(model)
    payload = model.model_dump(mode="json")
    return {"ok": True, "business_id": business_id, "appended": appended, "payload": payload, "receipt": _receipt(business_id=business_id, kind="business_map", payload=payload)}


@router.post("/business-map/{candidate}/governance/facts")
def propose_governed_business_map_fact(candidate: str, request: BusinessMapFactProposalRequest) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    try:
        return _business_map_governance().propose_fact(business_id, **request.model_dump())
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail="business_map_not_found") from exc
    except ValueError as exc:
        status = 409 if "revision_conflict" in str(exc) else 400
        raise HTTPException(status_code=status, detail=str(exc)) from exc


@router.post("/business-map/{candidate}/governance/facts/{fact_id}/review")
def review_governed_business_map_fact(
    candidate: str, fact_id: str, request: BusinessMapFactReviewRequest
) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    try:
        return _business_map_governance().review_fact(
            business_id, fact_id, **request.model_dump()
        )
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail="business_map_not_found") from exc
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc).strip("'")) from exc
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc
    except ValueError as exc:
        status = 409 if "revision_conflict" in str(exc) else 400
        raise HTTPException(status_code=status, detail=str(exc)) from exc


@router.post("/business-map/{candidate}/governance/facts/merge")
def merge_governed_business_map_facts(candidate: str, request: BusinessMapFactMergeRequest) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    try:
        return _business_map_governance().merge_facts(business_id, **request.model_dump())
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail="business_map_not_found") from exc
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc).strip("'")) from exc
    except ValueError as exc:
        status = 409 if "revision_conflict" in str(exc) else 400
        raise HTTPException(status_code=status, detail=str(exc)) from exc


@router.post("/business-map/{candidate}/governance/relationships")
def propose_governed_business_map_relationship(candidate: str, request: BusinessMapRelationshipProposalRequest) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    try:
        return _business_map_governance().propose_relationship(business_id, **request.model_dump())
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail="business_map_not_found") from exc
    except ValueError as exc:
        status = 409 if "revision_conflict" in str(exc) else 400
        raise HTTPException(status_code=status, detail=str(exc)) from exc


@router.post("/business-map/{candidate}/governance/relationships/{relationship_id}/review")
def review_governed_business_map_relationship(candidate: str, relationship_id: str, request: BusinessMapRelationshipReviewRequest) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    try:
        return _business_map_governance().review_relationship(business_id, relationship_id, **request.model_dump())
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail="business_map_not_found") from exc
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc).strip("'")) from exc
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc
    except ValueError as exc:
        status = 409 if "revision_conflict" in str(exc) else 400
        raise HTTPException(status_code=status, detail=str(exc)) from exc


@router.post("/business-map/{candidate}/governance/relationships/normalize-legacy")
def normalize_legacy_business_map_relationships(candidate: str, request: BusinessMapRelationshipNormalizeRequest) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    try:
        return _business_map_governance().normalize_legacy_relationships(business_id, **request.model_dump())
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail="business_map_not_found") from exc
    except ValueError as exc:
        status = 409 if "revision_conflict" in str(exc) else 400
        raise HTTPException(status_code=status, detail=str(exc)) from exc


@router.get("/business-map/{candidate}/governance/queues")
def get_business_map_governance_queues(candidate: str) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    try:
        return _business_map_governance().queues(business_id)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail="business_map_not_found") from exc


@router.get("/business-map/{candidate}/governance/export")
def export_governed_business_map(candidate: str) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    try:
        return _business_map_governance().export(business_id)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=404, detail="business_map_not_found") from exc


@router.get("/business-knowledge/{candidate}")
def get_business_knowledge(candidate: str) -> Dict[str, Any]:
    try:
        business_id = canonical_business_id(candidate)
        summary = _business_knowledge().summary(business_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True, "business_id": business_id, "payload": summary}


@router.post("/business-knowledge/{candidate}/text")
def ingest_business_knowledge_text(candidate: str, request: BusinessKnowledgeTextRequest) -> Dict[str, Any]:
    try:
        business_id = canonical_business_id(candidate)
        result = _business_knowledge().ingest_text(
            business_id,
            title=request.title,
            text=request.text,
            scopes=request.scopes,
            confidentiality=request.confidentiality,
            effective_from=request.effective_from,
            effective_until=request.effective_until,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True, "business_id": business_id, **result}


@router.post("/business-knowledge/{candidate}/file")
async def ingest_business_knowledge_file(
    candidate: str,
    file: UploadFile = File(...),
    scopes: str = Form("all"),
    confidentiality: str = Form("internal"),
    effective_from: str | None = Form(None),
    effective_until: str | None = Form(None),
) -> Dict[str, Any]:
    try:
        business_id = canonical_business_id(candidate)
        data = await file.read()
        scope_values = [value.strip() for value in scopes.split(",") if value.strip()]
        result = _business_knowledge().ingest_file(
            business_id,
            filename=file.filename or "business-knowledge.txt",
            data=data,
            scopes=scope_values,
            confidentiality=confidentiality,
            effective_from=effective_from,
            effective_until=effective_until,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True, "business_id": business_id, **result}


@router.post("/business-knowledge/{candidate}/claims/approve")
def approve_business_knowledge_claims(candidate: str, request: BusinessKnowledgeApprovalRequest) -> Dict[str, Any]:
    try:
        business_id = canonical_business_id(candidate)
        result = _business_knowledge().approve(business_id, request.claim_ids, actor=request.actor)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    approved = result.get("approved") or []
    if approved:
        append_business_map_facts(
            business_id,
            BusinessMapFactsRequest(
                source="business_knowledge_owner_approval",
                facts=[
                    {
                        "function": (claim.get("scopes") or ["all"])[0],
                        "category": "business_knowledge",
                        "field": claim.get("claim_id"),
                        "value": claim.get("text"),
                        "classification": "owner_attested_business_knowledge",
                        "confidence": 1.0,
                        "verification_status": "owner_attested_not_independently_verified",
                        "source_ref": f"business_knowledge:{claim.get('source_id')}",
                        "evidence_refs": [claim.get("source_hash")],
                        "confidentiality": claim.get("confidentiality"),
                        "scopes": claim.get("scopes"),
                        "effective_from": claim.get("effective_from"),
                        "effective_until": claim.get("effective_until"),
                    }
                    for claim in approved
                ],
            ),
        )
    return {"ok": True, "business_id": business_id, **result}


@router.post("/business-knowledge/{candidate}/search")
def search_business_knowledge(candidate: str, request: BusinessKnowledgeSearchRequest) -> Dict[str, Any]:
    try:
        business_id = canonical_business_id(candidate)
        result = _business_knowledge().search(
            business_id,
            request.query,
            actor_scope=request.actor_scope,
            limit=request.limit,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True, "business_id": business_id, **result}


@router.post("/business-knowledge/{candidate}/entity-links/rebuild")
def rebuild_business_knowledge_entity_links(candidate: str, request: BusinessKnowledgeEntityLinkRequest) -> Dict[str, Any]:
    try:
        business_id = canonical_business_id(candidate)
        result = _business_knowledge().rebuild_entity_links(business_id, actor=request.actor)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True, **result}


@router.post("/business-knowledge/{candidate}/sources/{source_id}/policy")
def set_business_knowledge_source_policy(candidate: str, source_id: str, request: BusinessKnowledgeSourcePolicyRequest) -> Dict[str, Any]:
    try:
        business_id = canonical_business_id(candidate)
        result = _business_knowledge().set_source_policy(business_id, source_id, **request.model_dump())
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc).strip("'")) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True, "business_id": business_id, **result}


@router.post("/business-knowledge/{candidate}/sources/{source_id}/revoke")
def revoke_business_knowledge_source(candidate: str, source_id: str, request: BusinessKnowledgeSourceRevokeRequest) -> Dict[str, Any]:
    try:
        business_id = canonical_business_id(candidate)
        result = _business_knowledge().revoke_source(business_id, source_id, **request.model_dump())
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc).strip("'")) from exc
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True, "business_id": business_id, **result}


@router.post("/business-knowledge/{candidate}/retention/apply")
def apply_business_knowledge_retention(candidate: str, request: BusinessKnowledgeRetentionRequest) -> Dict[str, Any]:
    try:
        business_id = canonical_business_id(candidate)
        result = _business_knowledge().apply_retention(business_id, actor=request.actor)
    except (PermissionError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return {"ok": True, **result}


@router.get("/department-intelligence/{candidate}")
def get_department_intelligence(candidate: str) -> Dict[str, Any]:
    return _read_container(candidate, "department_intelligence")


@router.put("/department-intelligence/{candidate}")
def put_department_intelligence(candidate: str, request: ContainerWriteRequest) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    repository = _repo()
    existing = repository.load_optional_dict(business_id, "department_intelligence")
    payload = _object(request.payload)
    payload["departments"] = {
        **_object((existing or {}).get("departments")),
        **_object(payload.get("departments")),
    }
    payload.update({"id": f"{business_id}.department_intelligence", "workspace_id": business_id, "kind": "department_intelligence"})
    payload["revision"] = _next_revision(existing, request.expected_revision)
    payload["meta"] = _meta(business_id, "department_intelligence", request.source).model_dump(mode="json")
    model = DepartmentIntelligenceContainer(**payload)
    repository.save_model(model)
    data = model.model_dump(mode="json")
    return {"ok": True, "business_id": business_id, "payload": data, "receipt": _receipt(business_id=business_id, kind="department_intelligence", payload=data)}


@router.get("/finance-model/{candidate}")
def get_finance_model(candidate: str) -> Dict[str, Any]:
    return _read_container(candidate, "business_financial_model")


@router.put("/finance-model/{candidate}")
def put_finance_model(candidate: str, request: ContainerWriteRequest) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    repository = _repo()
    existing = repository.load_optional_dict(business_id, "business_financial_model")
    payload = _object(request.payload)
    # Desktop discovery saves must not erase provider evidence or accepted
    # document facts already committed by authoritative backend workflows.
    if existing:
        if existing.get("external_data"):
            payload["external_data"] = _object(existing.get("external_data"))
        payload["integration_evidence"] = {
            **_object(existing.get("integration_evidence")),
            **_object(payload.get("integration_evidence")),
        }
        refs: list[Dict[str, Any]] = []
        seen: set[str] = set()
        for item in [*_list(existing.get("evidence_refs")), *_list(payload.get("evidence_refs"))]:
            if not isinstance(item, dict):
                continue
            key = str(item.get("artifact_id") or item.get("source") or _hash(item))
            if key in seen:
                continue
            seen.add(key)
            refs.append(item)
        payload["evidence_refs"] = refs
        # A resumed desktop discovery state must not replace a completed,
        # source-backed statement package with browser-parsed approximations.
        if existing.get("authoritative_metrics"):
            payload["metrics"] = _object(existing.get("authoritative_metrics"))
            for key in (
                "authoritative_metrics",
                "financial_statements",
                "business_model",
                "statement_quality",
                "reporting_period",
                "finance_director",
            ):
                payload[key] = _object(existing.get(key))
    payload.update({"id": f"{business_id}.business_financial_model", "workspace_id": business_id, "kind": "business_financial_model"})
    payload["revision"] = _next_revision(existing, request.expected_revision)
    payload["meta"] = _meta(business_id, "business_financial_model", request.source).model_dump(mode="json")
    model = BusinessFinancialModelContainer(**payload)
    repository.save_model(model)
    data = model.model_dump(mode="json")
    return {"ok": True, "business_id": business_id, "payload": data, "receipt": _receipt(business_id=business_id, kind="business_financial_model", payload=data)}


@router.post("/finance-completion/{candidate}")
def complete_finance_foundation(candidate: str) -> Dict[str, Any]:
    """Create source-backed statements, seed the operating model and update Boardroom.

    This is intentionally idempotent for user-owned operating data: an existing
    Products & Services model is never replaced by a newly inferred seed.
    """

    business_id = canonical_business_id(candidate)
    repository = _repo()
    existing_financial = repository.load_optional_dict(business_id, "business_financial_model")
    if not existing_financial:
        raise HTTPException(status_code=404, detail="business_financial_model_not_found")

    existing_operating = repository.load_optional_dict(business_id, "business_operating_model")
    initial_completion = build_finance_completion(existing_financial, existing_operating)

    operating_created = existing_operating is None
    operating_payload = existing_operating or seed_operating_model(existing_financial, initial_completion)
    operating_payload = _calculate_operating_model(operating_payload)
    if operating_created:
        operating_payload.update({
            "id": f"{business_id}.business_operating_model",
            "workspace_id": business_id,
            "kind": "business_operating_model",
            "revision": 1,
            "meta": _meta(business_id, "business_operating_model", "finance_completion").model_dump(mode="json"),
        })
        operating_model = BusinessOperatingModelContainer(**operating_payload)
        repository.save_model(operating_model)
        operating_payload = operating_model.model_dump(mode="json")
        _project_operating_model_to_business_map(
            repository, business_id, operating_payload, "finance_completion"
        )

    completion = build_finance_completion(existing_financial, operating_payload)
    financial_payload = {
        **existing_financial,
        "currency": completion["currency"],
        "period_basis": completion["period"]["basis"],
        "reporting_period": completion["period"],
        "metrics": completion["authoritative_metrics"],
        "authoritative_metrics": completion["authoritative_metrics"],
        "financial_statements": completion["financial_statements"],
        "business_model": completion["business_model"],
        "statement_quality": completion["quality"],
        "revision": int(existing_financial.get("revision") or 0) + 1,
        "meta": _meta(business_id, "business_financial_model", "finance_completion").model_dump(mode="json"),
        "provenance": {
            **_object(existing_financial.get("provenance")),
            "finance_completed_at": completion["generated_at"],
            "finance_completion_schema": completion["schema_version"],
            "evidence_precedence": [
                "accepted_accounting_evidence",
                "confirmed_founder_input",
                "calculated_from_source_backed_values",
                "unverified_discovery_input",
            ],
        },
    }
    financial_model = BusinessFinancialModelContainer(**financial_payload)
    repository.save_model(financial_model)
    financial_payload = financial_model.model_dump(mode="json")

    intelligence = repository.load_optional_dict(business_id, "department_intelligence") or {}
    departments = _object(intelligence.get("departments"))
    finance_department = _object(departments.get("finance"))
    finance_department["financial_statements"] = {
        "period": completion["period"],
        "profit_and_loss_status": completion["financial_statements"]["profit_and_loss"]["status"],
        "balance_sheet_status": completion["financial_statements"]["balance_sheet"]["status"],
        "cash_flow_status": completion["financial_statements"]["cash_flow"]["status"],
        "quality": completion["quality"],
        "financial_model_revision": financial_payload["revision"],
        "updated_at": completion["generated_at"],
    }
    finance_department["operating_model"] = {
        "status": operating_payload.get("model_status"),
        "operating_model_revision": operating_payload.get("revision"),
        "offering_count": len(_list(operating_payload.get("offerings"))),
        "requires_owner_review": operating_payload.get("model_status") == "needs_user_review",
        "updated_at": completion["generated_at"],
    }
    departments["finance"] = finance_department
    departments["boardroom"] = {
        **_object(departments.get("boardroom")),
        "status": "finance_statements_ready_for_review",
        "finance_completion": {
            "period": completion["period"],
            "quality_status": completion["quality"]["status"],
            "financial_model_revision": financial_payload["revision"],
            "operating_model_revision": operating_payload.get("revision"),
            "updated_at": completion["generated_at"],
        },
    }
    intelligence_payload = {
        **intelligence,
        "id": f"{business_id}.department_intelligence",
        "workspace_id": business_id,
        "kind": "department_intelligence",
        "meta": _meta(business_id, "department_intelligence", "finance_completion").model_dump(mode="json"),
        "departments": departments,
        "revision": int(intelligence.get("revision") or 0) + 1,
    }
    repository.save_model(DepartmentIntelligenceContainer(**intelligence_payload))

    # The Director refresh writes its source-linked projection into the
    # financial model, Department Intelligence and the live Boardroom runtime.
    director = FinanceDirectorService(repository).refresh(
        business_id,
        minimum_cash_reserve=completion["financial_statements"]["cash_flow"].get("protected_cash_reserve"),
        created_by="finance_completion",
        created_at=completion["generated_at"],
    )

    return {
        "ok": True,
        "business_id": business_id,
        "operating_model_created": operating_created,
        "financial_model": repository.load_dict(business_id, "business_financial_model"),
        "operating_model": operating_payload,
        "finance_director": director,
        "boardroom_projection_updated": True,
        "external_action_performed": False,
    }


@router.get("/operating-model/{candidate}")
def get_operating_model(candidate: str) -> Dict[str, Any]:
    return _read_container(candidate, "business_operating_model")


@router.put("/operating-model/{candidate}")
def put_operating_model(candidate: str, request: ContainerWriteRequest) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    repository = _repo()
    existing = repository.load_optional_dict(business_id, "business_operating_model")
    source_payload = _object(request.payload)
    organisation = repository.load_optional_dict(business_id, "organization_authority") or {}
    people = {str(item.get("id")): item for item in _list(organisation.get("people")) if isinstance(item, dict)}
    hydrated_labour = []
    for raw in _list(source_payload.get("labour_resources")):
        resource = _object(raw)
        person = people.get(str(resource.get("person_id") or ""))
        if person:
            costing = _object(person.get("workforce_costing"))
            requested_basis = str(resource.get("cost_basis") or costing.get("basis") or "hourly")
            manual_override = (
                str(resource.get("rate_source") or "") == "manual_planning_override"
                and resource.get("cost_rate") not in (None, "")
            )
            if manual_override:
                effective_rate = resource.get("cost_rate")
                cost_source = "owner_planning_override"
            elif requested_basis == "daily":
                effective_rate = costing.get("effective_daily_cost")
                cost_source = "confidential_hr_effective_planning_cost"
            elif requested_basis == "monthly":
                effective_rate = costing.get("effective_monthly_cost")
                cost_source = "confidential_hr_effective_planning_cost"
            elif requested_basis == "per_job":
                effective_rate = costing.get("base_rate")
                cost_source = "confidential_hr_effective_planning_cost"
            else:
                requested_basis = "hourly"
                effective_rate = costing.get("effective_hourly_cost")
                cost_source = "confidential_hr_effective_planning_cost"
            resource.update({
                "person_name": person.get("name"), "engagement_type": person.get("employment_type"),
                "cost_basis": requested_basis, "cost_rate": effective_rate,
                "cost_source": cost_source,
            })
        hydrated_labour.append(resource)
    source_payload["labour_resources"] = hydrated_labour
    payload = _calculate_operating_model(source_payload)
    payload.update({
        "id": f"{business_id}.business_operating_model",
        "workspace_id": business_id,
        "kind": "business_operating_model",
    })
    payload["revision"] = _next_revision(existing, request.expected_revision)
    payload["meta"] = _meta(business_id, "business_operating_model", request.source).model_dump(mode="json")
    payload["provenance"] = {
        **_object(payload.get("provenance")),
        "last_saved_at": _now(),
        "last_saved_by": request.source,
        "calculated_by": "aion_operating_model.v1",
    }
    model = BusinessOperatingModelContainer(**payload)
    repository.save_model(model)
    data = model.model_dump(mode="json")
    _project_operating_model_to_business_map(repository, business_id, data, request.source)

    ledger = repository.load_optional_dict(business_id, "department_intelligence") or {}
    departments = _object(ledger.get("departments"))
    summary = {
        "status": data.get("model_status"),
        "operating_model_revision": data.get("revision"),
        "offering_count": len(_list(data.get("offerings"))),
        "inventory_item_count": len(_list(data.get("inventory_items"))),
        "stock_value": _object(data.get("inventory_metrics")).get("total_stock_value"),
        "procurement_items_requiring_review": len([
            item for item in _list(data.get("procurement_queue"))
            if str(_object(item).get("status") or "") in {"suggested", "draft"}
        ]),
        "updated_at": _now(),
    }
    for department in ("products_services", "operations", "sales", "finance", "hr", "procurement"):
        departments[department] = {**_object(departments.get(department)), "department": department, "operating_model": summary}
    ledger_payload = {
        **ledger,
        "id": f"{business_id}.department_intelligence",
        "workspace_id": business_id,
        "kind": "department_intelligence",
        "meta": _meta(business_id, "department_intelligence", request.source).model_dump(mode="json"),
        "departments": departments,
        "revision": int(ledger.get("revision") or 0) + 1,
    }
    repository.save_model(DepartmentIntelligenceContainer(**ledger_payload))
    return {
        "ok": True, "business_id": business_id, "payload": data,
        "receipt": _receipt(business_id=business_id, kind="business_operating_model", payload=data),
        "boardroom_projection_updated": True,
    }


@router.post("/operating-model/{candidate}/import")
async def import_operating_model_spreadsheet(
    candidate: str,
    category: str = Form(...),
    file: UploadFile = File(...),
) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    if category not in {"offerings", "inventory"}:
        raise HTTPException(status_code=400, detail="operating_model_import_category_invalid")
    filename = _safe_filename(file.filename or f"{category}.xlsx")
    extension = Path(filename).suffix.lower()
    content = await file.read()
    if not content or len(content) > MAX_FINANCE_ARTIFACT_BYTES:
        raise HTTPException(status_code=400, detail="operating_model_import_empty_or_too_large")
    rows = _map_operating_import_rows(_parse_operating_model_rows(content, extension), category)
    if not rows:
        raise HTTPException(status_code=422, detail="no_recognised_operating_model_rows")

    artifact_hash = "sha256:" + sha256(content).hexdigest()
    artifact_id = f"operating-import-{sha256((artifact_hash + category).encode()).hexdigest()[:16]}"
    storage_dir = AIONBusinessPaths.business_container_dir(business_id) / "artifacts" / "operating_model" / category
    storage_dir.mkdir(parents=True, exist_ok=True)
    stored_name = f"{artifact_id}{extension}"
    storage_path = storage_dir / stored_name
    storage_path.write_bytes(content)
    relative_path = str(storage_path.relative_to(AIONBusinessPaths.ROOT))

    repository = _repo()
    existing = repository.load_optional_dict(business_id, "business_operating_model") or {}
    target_key = "offerings" if category == "offerings" else "inventory_items"
    existing[target_key] = [*_list(existing.get(target_key)), *rows]
    source_refs = _list(existing.get("source_refs"))
    record = {
        "artifact_id": artifact_id, "name": filename, "extension": extension.lstrip("."),
        "artifact_hash": artifact_hash, "storage_path": relative_path,
        "imported_row_count": len(rows), "category": category, "imported_at": _now(),
    }
    source_refs.append(record)
    existing["source_refs"] = source_refs
    saved = put_operating_model(
        business_id,
        ContainerWriteRequest(payload=existing, source="operating_model_spreadsheet_import", expected_revision=int(existing.get("revision") or 0) or None),
    )
    pointer = _add_operating_model_file_pointer(business_id=business_id, category=category, record=record)
    return {**saved, "import": record, "imported_rows": rows, "file_cabinet_pointer": pointer}


@router.post("/finance-artifacts/{candidate}")
async def upload_finance_artifact(
    candidate: str,
    category: str = Form(...),
    file: UploadFile = File(...),
) -> Dict[str, Any]:
    try:
        business_id = canonical_business_id(candidate)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    category = str(category or "").strip().lower().replace("_", "-")
    allowed = FINANCE_ARTIFACT_CATEGORIES.get(category)
    if allowed is None:
        raise HTTPException(status_code=400, detail="unsupported_finance_artifact_category")

    filename = _safe_filename(file.filename or "upload")
    extension = Path(filename).suffix.lower()
    if extension not in allowed:
        raise HTTPException(status_code=400, detail=f"unsupported_file_type_for_{category}:{extension or 'none'}")

    data = await file.read(MAX_FINANCE_ARTIFACT_BYTES + 1)
    if not data:
        raise HTTPException(status_code=400, detail="empty_file")
    if len(data) > MAX_FINANCE_ARTIFACT_BYTES:
        raise HTTPException(status_code=413, detail="finance_artifact_too_large")

    digest = sha256(data).hexdigest()
    artifact_id = f"artifact_{digest[:24]}"
    artifact_dir = AIONBusinessPaths.business_container_dir(business_id) / "finance" / category
    artifact_dir.mkdir(parents=True, exist_ok=True)
    stored_name = f"{artifact_id}_{filename}"
    destination = (artifact_dir / stored_name).resolve()
    business_root = AIONBusinessPaths.business_container_dir(business_id).resolve()
    try:
        destination.relative_to(business_root)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="artifact_path_escape") from exc

    destination.write_bytes(data)
    relative_storage_path = f"business_containers/{business_id}/finance/{category}/{stored_name}"
    target = {
        "business_id": business_id,
        "business_container_id": business_id,
        "sub_container": "finance",
        "artifact_type": "spreadsheet" if extension in {".csv", ".xlsx", ".xls"} else "document",
        "artifact_name": filename,
    }
    provenance = {
        "mission_id": "finance_onboarding",
        "mission_run_id": f"finance_upload_{artifact_id}",
        "step_id": f"finance_{category}_upload",
        "tool_id": "finance_file_ingestion",
        "created_by": "user_upload",
        "uploaded_at": _now(),
    }
    source_payload_hash = _hash({"filename": filename, "size": len(data), "mime_type": file.content_type, "category": category})
    receipt_hash = _hash({"artifact_hash": f"sha256:{digest}", "source_payload_hash": source_payload_hash, "provenance": provenance, "storage_path": relative_storage_path})
    record = {
        "schema_version": "aion.business_container_artifact_record.v1",
        "artifact_id": artifact_id,
        "target": target,
        "provenance": provenance,
        "artifact_hash": f"sha256:{digest}",
        "source_payload_hash": source_payload_hash,
        "receipt_hash": receipt_hash,
        "storage_path": relative_storage_path,
        "storage_state": "committed_to_business_container",
        "session_vfs_only": False,
        "live_side_effects_enabled": False,
        "byte_size": len(data),
        "mime_type": file.content_type or "application/octet-stream",
        "category": category,
    }
    record["replay_locator"] = {
        "business_id": business_id,
        "business_container_id": business_id,
        "department_id": "finance",
        "category": category,
        "storage_path": relative_storage_path,
        "artifact_id": artifact_id,
        "receipt_hash": receipt_hash,
    }
    record["replay_locator_hash"] = _hash(record["replay_locator"])
    receipt_path = artifact_dir / f"{artifact_id}.artifact.json"
    receipt_path.write_text(json.dumps(record, indent=2, sort_keys=True), encoding="utf-8")

    pointer = _add_finance_file_cabinet_pointer(business_id=business_id, category=category, record=record)
    return {"ok": True, "business_id": business_id, "record": record, "file_cabinet_pointer": pointer}


def _finance_artifact_record(business_id: str, artifact_id: str) -> tuple[Path, Dict[str, Any]]:
    finance_root = AIONBusinessPaths.business_container_dir(business_id) / "finance"
    matches = list(finance_root.glob(f"*/{artifact_id}.artifact.json"))
    if not matches:
        raise HTTPException(status_code=404, detail="finance_artifact_not_found")
    receipt_path = matches[0]
    record = _object(json.loads(receipt_path.read_text(encoding="utf-8")))
    if record.get("business_id") not in {None, business_id} and _object(record.get("target")).get("business_id") != business_id:
        raise HTTPException(status_code=409, detail="finance_artifact_business_mismatch")
    return receipt_path, record


@router.get("/finance-artifacts/{candidate}")
def list_finance_artifacts(candidate: str) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    finance_root = AIONBusinessPaths.business_container_dir(business_id) / "finance"
    artifacts: list[Dict[str, Any]] = []
    for receipt_path in sorted(finance_root.glob("*/*.artifact.json")):
        record = _object(json.loads(receipt_path.read_text(encoding="utf-8")))
        artifact_id = str(record.get("artifact_id") or "")
        analysis_path = receipt_path.parent / f"{artifact_id}.analysis.json"
        artifacts.append({
            "record": record,
            "analysis": json.loads(analysis_path.read_text(encoding="utf-8")) if analysis_path.exists() else None,
        })
    return {"ok": True, "business_id": business_id, "artifacts": artifacts}


def _finance_artifact_bytes_path(business_id: str, record: Dict[str, Any]) -> Path:
    relative = str(record.get("storage_path") or "")
    path = (AIONBusinessPaths.ROOT / relative).resolve()
    root = AIONBusinessPaths.business_container_dir(business_id).resolve()
    try:
        path.relative_to(root)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="artifact_path_escape") from exc
    if not path.exists():
        raise HTTPException(status_code=404, detail="finance_artifact_bytes_missing")
    return path


@router.post("/finance-artifacts/{candidate}/{artifact_id}/analyse")
def analyse_uploaded_finance_artifact(candidate: str, artifact_id: str) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    receipt_path, record = _finance_artifact_record(business_id, artifact_id)
    source_path = _finance_artifact_bytes_path(business_id, record)
    try:
        analysis = analyse_finance_artifact(
            source_path,
            filename=str(_object(record.get("target")).get("artifact_name") or source_path.name),
            category=str(record.get("category") or receipt_path.parent.name),
            artifact_id=artifact_id,
        )
    except (ValueError, json.JSONDecodeError) as exc:
        raise HTTPException(status_code=422, detail=f"finance_artifact_parse_failed:{exc}") from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc

    analysis_path = receipt_path.parent / f"{artifact_id}.analysis.json"
    analysis_path.write_text(json.dumps(analysis, indent=2, sort_keys=True, default=str), encoding="utf-8")
    analysis_ref = {
        "status": analysis["status"],
        "analysis_id": analysis["analysis_id"],
        "analysis_hash": analysis["analysis_hash"],
        "storage_path": f"business_containers/{business_id}/finance/{record['category']}/{analysis_path.name}",
        "candidate_fact_count": len(analysis.get("candidate_facts") or []),
        "acceptance_required": True,
        "analysed_at": analysis["analysed_at"],
    }
    record["analysis"] = analysis_ref
    receipt_path.write_text(json.dumps(record, indent=2, sort_keys=True), encoding="utf-8")
    pointer = _add_finance_file_cabinet_pointer(
        business_id=business_id,
        category=str(record["category"]),
        record=record,
    )
    return {
        "ok": True,
        "business_id": business_id,
        "artifact_id": artifact_id,
        "analysis": analysis,
        "analysis_ref": analysis_ref,
        "file_cabinet_pointer": pointer,
    }


@router.get("/finance-artifacts/{candidate}/{artifact_id}/analysis")
def get_finance_artifact_analysis(candidate: str, artifact_id: str) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    receipt_path, _record = _finance_artifact_record(business_id, artifact_id)
    analysis_path = receipt_path.parent / f"{artifact_id}.analysis.json"
    if not analysis_path.exists():
        raise HTTPException(status_code=404, detail="finance_artifact_not_analysed")
    return {"ok": True, "business_id": business_id, "artifact_id": artifact_id, "analysis": json.loads(analysis_path.read_text(encoding="utf-8"))}


@router.post("/finance-artifacts/{candidate}/{artifact_id}/accept")
def accept_finance_artifact_facts(
    candidate: str,
    artifact_id: str,
    request: FinanceArtifactAcceptanceRequest,
) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    receipt_path, record = _finance_artifact_record(business_id, artifact_id)
    analysis_path = receipt_path.parent / f"{artifact_id}.analysis.json"
    if not analysis_path.exists():
        raise HTTPException(status_code=409, detail="analyse_artifact_before_accepting_facts")
    analysis = _object(json.loads(analysis_path.read_text(encoding="utf-8")))
    candidates = [item for item in _list(analysis.get("candidate_facts")) if isinstance(item, dict)]
    requested = set(request.fact_ids or [str(item.get("fact_id") or "") for item in candidates])
    accepted = [item for item in candidates if str(item.get("fact_id") or "") in requested]
    if not accepted:
        raise HTTPException(status_code=400, detail="no_candidate_facts_selected")

    repository = _repo()
    try:
        business_map = repository.load_business_map(business_id)
    except FileNotFoundError:
        business_map = BusinessMapContainer(
            id=f"{business_id}.business_map",
            workspace_id=business_id,
            meta=_meta(business_id, "business_map", "finance_artifact_acceptance"),
        )
    existing_ids = {str(item.get("fact_id") or "") for item in business_map.facts if isinstance(item, dict)}
    observed_at = _now()
    appended = []
    for candidate_fact in accepted:
        fact = {
            **candidate_fact,
            "business_id": business_id,
            "classification": "user_accepted_extracted_fact",
            "verification_status": "accepted_from_source_document",
            "accepted": True,
            "accepted_at": observed_at,
            "source_ref": f"finance_artifact:{artifact_id}:{analysis.get('analysis_hash')}",
            "evidence_refs": [{
                "artifact_id": artifact_id,
                "artifact_hash": record.get("artifact_hash"),
                "analysis_hash": analysis.get("analysis_hash"),
            }],
            "observed_at": observed_at,
        }
        if str(fact.get("fact_id") or "") not in existing_ids:
            business_map.facts.append(fact)
            appended.append(fact)
            existing_ids.add(str(fact.get("fact_id") or ""))
    if appended:
        business_map.revision = int(business_map.revision or 0) + 1
        business_map.meta.updated_at = observed_at
        business_map.meta.source = "finance_artifact_acceptance"
        repository.save_model(business_map)

    finance_payload = repository.load_optional_dict(business_id, "business_financial_model") or {
        "id": f"{business_id}.business_financial_model",
        "workspace_id": business_id,
        "kind": "business_financial_model",
        "meta": _meta(business_id, "business_financial_model", "finance_artifact_acceptance").model_dump(mode="json"),
    }
    refs = _list(finance_payload.get("evidence_refs"))
    refs = [item for item in refs if item.get("artifact_id") != artifact_id]
    refs.append({
        "source": "user_uploaded_finance_artifact",
        "artifact_id": artifact_id,
        "artifact_hash": record.get("artifact_hash"),
        "analysis_hash": analysis.get("analysis_hash"),
        "accepted_fact_ids": [item.get("fact_id") for item in accepted],
        "storage_path": record.get("storage_path"),
        "verification_status": "accepted_from_source_document",
        "accepted_at": observed_at,
    })
    finance_payload["evidence_refs"] = refs
    external_data = _object(finance_payload.get("external_data"))
    artifacts = _object(external_data.get("artifacts"))
    artifacts[artifact_id] = {
        "artifact_id": artifact_id,
        "filename": _object(record.get("target")).get("artifact_name"),
        "category": record.get("category"),
        "artifact_hash": record.get("artifact_hash"),
        "analysis_hash": analysis.get("analysis_hash"),
        "verification_status": "accepted_from_source_document",
        "accepted_at": observed_at,
        "facts": [{
            "fact_id": item.get("fact_id"),
            "field": item.get("field"),
            "value": item.get("value"),
            "unit": item.get("unit"),
            "aggregation": item.get("aggregation"),
            "source_sheet": item.get("source_sheet"),
            "source_column": item.get("source_column"),
        } for item in accepted],
    }
    external_data["artifacts"] = artifacts
    external_data["accepted_artifact_count"] = len(artifacts)
    external_data["accepted_fact_count"] = sum(
        len(_list(_object(item).get("facts"))) for item in artifacts.values()
    )
    external_data["updated_at"] = observed_at
    finance_payload["external_data"] = external_data
    finance_payload["revision"] = int(finance_payload.get("revision") or 0) + 1
    finance_payload["meta"] = _meta(business_id, "business_financial_model", "finance_artifact_acceptance").model_dump(mode="json")
    repository.save_model(BusinessFinancialModelContainer(**finance_payload))

    accepted_ids = [str(item.get("fact_id") or "") for item in accepted]
    analysis["status"] = "accepted_into_finance_model"
    analysis["accepted_fact_ids"] = accepted_ids
    analysis["accepted_at"] = observed_at
    for item in candidates:
        item["accepted"] = str(item.get("fact_id") or "") in set(accepted_ids)
    analysis["candidate_facts"] = candidates
    analysis_path.write_text(json.dumps(analysis, indent=2, sort_keys=True, default=str), encoding="utf-8")
    record["analysis"] = {
        **_object(record.get("analysis")),
        "status": analysis["status"],
        "accepted_fact_ids": accepted_ids,
        "accepted_at": observed_at,
    }
    receipt_path.write_text(json.dumps(record, indent=2, sort_keys=True), encoding="utf-8")
    pointer = _add_finance_file_cabinet_pointer(business_id=business_id, category=str(record["category"]), record=record)
    return {
        "ok": True,
        "business_id": business_id,
        "artifact_id": artifact_id,
        "accepted_fact_count": len(accepted),
        "appended_fact_count": len(appended),
        "accepted_fact_ids": accepted_ids,
        "file_cabinet_pointer": pointer,
    }


@router.get("/boardroom-context/{candidate}")
def get_boardroom_context(candidate: str) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    repository = _repo()

    def optional(kind: str) -> Dict[str, Any]:
        return repository.load_optional_dict(business_id, kind) or {}

    business_map = optional("business_map")
    boardroom_facts = []
    for fact in _list(business_map.get("facts")):
        if not isinstance(fact, dict):
            continue
        if fact.get("classification") == "owner_attested_business_knowledge":
            scopes = {str(value) for value in _list(fact.get("scopes")) or ["all"]}
            if "all" not in scopes and "boardroom" not in scopes:
                continue
        boardroom_facts.append(fact)
    ledger = optional("department_intelligence")
    finance_model = optional("business_financial_model")
    operating_model = optional("business_operating_model")
    boardroom_operating_model = _boardroom_operating_projection(operating_model)
    evidence_refs = _list(finance_model.get("evidence_refs"))
    department_intelligence = _runtime_department_projection(
        business_id,
        _object(ledger.get("departments")),
        business_root=repository.base_dir / business_id,
    )
    envelope = {
        "schema_version": "aion.boardroom_context_envelope.v1",
        "business_id": business_id,
        "snapshot_id": f"boardroom_context_{business_id}_{int(datetime.now(UTC).timestamp())}",
        "generated_at": _now(),
        "persistent_truth_source": "business_containers",
        "boardroom_projection_only": True,
        "business_identity": optional("business_identity"),
        "business_structure": optional("business_structure"),
        "brand_foundation": optional("brand_foundation"),
        "business_map_projection": {
            "facts": boardroom_facts,
            "assumptions": _list(business_map.get("assumptions")),
            "unknowns": _list(business_map.get("unknowns")),
            "conflicts": _list(business_map.get("conflicts")),
            "revision": business_map.get("revision"),
        },
        "function_models": {"finance": finance_model, "operating_model": boardroom_operating_model},
        "operating_model": boardroom_operating_model,
        "department_intelligence": department_intelligence,
        "evidence_index": evidence_refs,
        "source_revisions": {
            "business_map": business_map.get("revision"),
            "department_intelligence": ledger.get("revision"),
            "finance": finance_model.get("revision"),
            "operating_model": operating_model.get("revision"),
        },
        "execution_boundary": {
            "approval_gated": True,
            "live_external_side_effects_performed": False,
            "raw_model_tool_access_allowed": False,
            "provider_memory_mutation_allowed": False,
        },
    }
    envelope["context_hash"] = _hash(envelope)
    return {"ok": True, "business_id": business_id, "context": envelope}


def _finance_grounded_response(question: str, model: Dict[str, Any]) -> str:
    metrics = _object(model.get("metrics"))
    missing = _list(model.get("missing_information"))
    lowered = question.lower()
    labels = {
        "average_monthly_revenue": "Average monthly revenue",
        "annualised_revenue": "Annualised revenue",
        "monthly_gross_profit": "Monthly gross profit",
        "gross_margin_percent": "Gross margin",
        "monthly_operating_surplus": "Monthly operating surplus",
        "operating_margin_percent": "Operating margin",
    }
    wanted = list(metrics)
    if "margin" in lowered:
        wanted = [key for key in metrics if "margin" in key or "gross_profit" in key]
    elif "cash" in lowered:
        cash = _object(model.get("cashflow_model"))
        return "Cash position on record: " + str(cash.get("cash_position") or "not yet quantified") + (f". Still needed: {', '.join(missing)}." if missing else ".")
    elif "fixed" in lowered or "cost" in lowered:
        return "Fixed-cost information on record: " + str(_object(model.get("overhead_model")).get("fixed_costs") or "not yet quantified") + ". Direct-cost information: " + str(_object(model.get("direct_cost_model")).get("direct_costs") or "not yet quantified") + ". Values remain unverified until reconciled to evidence."
    lines = []
    for key in wanted:
        item = _object(metrics.get(key))
        value = item.get("value")
        if value is None:
            continue
        suffix = "%" if "percent" in key else ""
        lines.append(f"{labels.get(key, key.replace('_', ' ').title())}: {value:,.2f}{suffix}")
    if not lines:
        return "The Finance Pilot does not yet have enough verified numeric data to answer that reliably. " + ("Still needed: " + ", ".join(missing) + "." if missing else "Complete or update Finance discovery first.")
    note = " These figures are calculated from founder-provided, currently unverified inputs."
    if missing:
        note += " Missing for a fuller model: " + ", ".join(missing) + "."
    return "Finance update — " + "; ".join(lines) + "." + note


@router.post("/finance-agent/{candidate}/turn")
def finance_agent_turn(candidate: str, request: FinanceAgentTurnRequest) -> Dict[str, Any]:
    business_id = canonical_business_id(candidate)
    repository = _repo()
    model = repository.load_optional_dict(business_id, "business_financial_model") or {}
    if not model:
        raise HTTPException(status_code=409, detail="finance_model_missing")
    response = _finance_grounded_response(request.user_text, model)
    transcript = _list(_object(model.get("discovery_state")).get("operational_transcript"))
    turn = {"user_text": request.user_text.strip(), "response": response, "created_at": _now(), "source": "canonical_finance_model", "live_external_action": False}
    transcript.append(turn)
    discovery = _object(model.get("discovery_state"))
    discovery["operational_transcript"] = transcript[-200:]
    model["discovery_state"] = discovery
    model["revision"] = int(model.get("revision") or 0) + 1
    model["meta"] = _meta(business_id, "business_financial_model", "finance_operational_terminal").model_dump(mode="json")
    repository.save_model(BusinessFinancialModelContainer(**model))
    return {"ok": True, "business_id": business_id, "turn": turn, "model_revision": model["revision"], "approval_gated": True}
