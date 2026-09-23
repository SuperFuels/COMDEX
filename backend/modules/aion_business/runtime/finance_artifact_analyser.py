from __future__ import annotations

import csv
from hashlib import sha256
import json
from pathlib import Path
import re
from statistics import fmean
from typing import Any, Dict, Iterable
import xml.etree.ElementTree as ET


MAX_ROWS = 5000
MAX_SAMPLE_ROWS = 20
MAX_PDF_TEXT = 12000


def _now_iso() -> str:
    from datetime import UTC, datetime
    return datetime.now(UTC).replace(microsecond=0).isoformat()


def _safe_number(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value or "").strip()
    if not text:
        return None
    text = re.sub(r"[^0-9,.()\-]", "", text)
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    if text.count(",") == 1 and "." not in text:
        left, right = text.split(",")
        text = left + ("." if len(right) <= 2 else "") + right
    else:
        text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _normalise_header(value: Any, index: int) -> str:
    text = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")
    return text or f"column_{index + 1}"


def _rows_analysis(rows: Iterable[Iterable[Any]], source_label: str) -> Dict[str, Any]:
    materialized = [list(row) for _, row in zip(range(MAX_ROWS + 1), rows)]
    truncated = len(materialized) > MAX_ROWS
    materialized = materialized[:MAX_ROWS]
    if not materialized:
        return {"source": source_label, "row_count": 0, "columns": [], "sample_rows": [], "numeric_columns": {}}
    headers = [_normalise_header(value, index) for index, value in enumerate(materialized[0])]
    data_rows = materialized[1:]
    numeric: Dict[str, list[float]] = {header: [] for header in headers}
    reported_totals: Dict[str, float] = {}
    samples: list[Dict[str, Any]] = []
    for row_index, row in enumerate(data_rows):
        mapped: Dict[str, Any] = {}
        row_label = str(row[0] if row else "").strip().lower()
        is_reported_total = bool(re.search(r"\b(total|closing)\b", row_label))
        for index, header in enumerate(headers):
            value = row[index] if index < len(row) else None
            mapped[header] = value
            number = _safe_number(value)
            if number is not None:
                if is_reported_total:
                    reported_totals[header] = number
                else:
                    numeric[header].append(number)
        if row_index < MAX_SAMPLE_ROWS:
            samples.append(mapped)
    summaries = {
        header: {
            "count": len(values),
            "sum": round(sum(values), 4),
            "average": round(fmean(values), 4),
            "minimum": round(min(values), 4),
            "maximum": round(max(values), 4),
        }
        for header, values in numeric.items() if values
    }
    return {
        "source": source_label,
        "row_count": len(data_rows),
        "row_limit_reached": truncated,
        "columns": headers,
        "sample_rows": samples,
        "numeric_columns": summaries,
        "reported_totals": reported_totals,
    }


def _candidate_facts(tabular: list[Dict[str, Any]]) -> list[Dict[str, Any]]:
    candidates: list[Dict[str, Any]] = []
    patterns = {
        "revenue": ("revenue", "sales", "income", "turnover", "net_sales", "gross_sales"),
        "costs": ("cost", "costs", "expense", "expenses", "cogs", "overheads"),
        "tax": ("tax", "vat", "gst", "sales_tax"),
        "cash": ("cash", "balance", "bank_balance"),
        "profit": ("profit", "net_profit", "gross_profit"),
    }
    for table in tabular:
        first_column = (table.get("columns") or [None])[0]
        for column, summary in (table.get("numeric_columns") or {}).items():
            # The first column normally contains month/row labels. Decorative
            # report titles can contain words such as "profit" or "cost" and
            # must never be promoted as monetary facts.
            if column == first_column:
                continue
            category = next((kind for kind, names in patterns.items() if column in names or any(name in column for name in names)), None)
            if not category:
                continue
            reported_total = (table.get("reported_totals") or {}).get(column)
            value = reported_total if reported_total is not None else summary.get("sum")
            aggregation = "reported_total_row" if reported_total is not None else "sum_of_detected_column"
            fact_id = "fact_" + sha256(f"{table.get('source')}:{column}:{value}".encode()).hexdigest()[:16]
            candidates.append({
                "fact_id": fact_id,
                "function": "finance",
                "field": f"external_data.{category}.{column}",
                "value": value,
                "unit": "source_currency",
                "aggregation": aggregation,
                "source_sheet": table.get("source"),
                "source_column": column,
                "classification": "extracted_candidate_fact",
                "confidence": 0.62,
                "verification_status": "requires_user_acceptance",
                "accepted": False,
            })
    return candidates


def analyse_finance_artifact(path: Path, *, filename: str, category: str, artifact_id: str) -> Dict[str, Any]:
    extension = path.suffix.lower()
    tables: list[Dict[str, Any]] = []
    document: Dict[str, Any] = {}
    warnings: list[str] = []

    if extension == ".csv":
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            tables.append(_rows_analysis(csv.reader(handle), filename))
    elif extension == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("openpyxl_dependency_missing") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            tables.append(_rows_analysis(sheet.iter_rows(values_only=True), sheet.title))
        workbook.close()
    elif extension == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("xlrd_dependency_missing") from exc
        workbook = xlrd.open_workbook(path, on_demand=True)
        for sheet in workbook.sheets():
            tables.append(_rows_analysis((sheet.row_values(index) for index in range(sheet.nrows)), sheet.name))
        workbook.release_resources()
    elif extension == ".json":
        value = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(value, list) and value and isinstance(value[0], dict):
            headers = list(dict.fromkeys(key for item in value[:MAX_ROWS] for key in item.keys()))
            tables.append(_rows_analysis([headers] + [[item.get(key) for key in headers] for item in value[:MAX_ROWS]], filename))
        else:
            document = {"json_type": type(value).__name__, "top_level_keys": list(value.keys())[:100] if isinstance(value, dict) else []}
    elif extension == ".xml":
        root = ET.parse(path).getroot()
        children = list(root)
        if children and all(list(child) for child in children[:10]):
            headers = list(dict.fromkeys(grand.tag.split("}")[-1] for child in children[:MAX_ROWS] for grand in list(child)))
            rows = [headers] + [[next((grand.text for grand in list(child) if grand.tag.split("}")[-1] == header), None) for header in headers] for child in children[:MAX_ROWS]]
            tables.append(_rows_analysis(rows, root.tag.split("}")[-1]))
        document = {"root_element": root.tag.split("}")[-1], "child_count": len(children)}
    elif extension == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("pypdf_dependency_missing") from exc
        reader = PdfReader(str(path))
        parts: list[str] = []
        for page in reader.pages[:100]:
            if sum(len(item) for item in parts) >= MAX_PDF_TEXT:
                break
            parts.append(page.extract_text() or "")
        text = "\n".join(parts)[:MAX_PDF_TEXT]
        document = {"page_count": len(reader.pages), "text_extract": text, "text_length": len(text)}
        if not text.strip():
            warnings.append("No machine-readable PDF text was found; OCR is not yet enabled.")
    else:
        document = {"format": extension.lstrip("."), "analysis": "metadata_only"}
        warnings.append("This format currently supports metadata analysis only.")

    analysis = {
        "schema_version": "aion.finance_artifact_analysis.v1",
        "analysis_id": f"analysis_{artifact_id}",
        "artifact_id": artifact_id,
        "filename": filename,
        "category": category,
        "analysed_at": _now_iso(),
        "status": "analysed_requires_review",
        "tables": tables,
        "document": document,
        "candidate_facts": _candidate_facts(tables),
        "warnings": warnings,
        "acceptance_required": True,
        "accepted_fact_ids": [],
    }
    analysis["analysis_hash"] = "sha256:" + sha256(
        json.dumps(analysis, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")
    ).hexdigest()
    return analysis
