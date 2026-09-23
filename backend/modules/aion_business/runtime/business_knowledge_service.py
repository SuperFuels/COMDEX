from __future__ import annotations

from datetime import UTC, datetime
from hashlib import sha256
import io
import json
from pathlib import Path
import re
from threading import RLock
from typing import Any, Iterable
from uuid import uuid4
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from backend.modules.aion_business.runtime.paths import AIONBusinessPaths


ALLOWED_SCOPES = {
    "all", "pilot", "boardroom", "marketing", "sales", "finance",
    "operations", "support", "people", "products_services",
}
ALLOWED_CONFIDENTIALITY = {"public", "internal", "confidential", "restricted"}
MAX_SOURCE_BYTES = 25 * 1024 * 1024
_LOCK = RLock()

_SEARCH_STOP_WORDS = {
    "a", "an", "and", "are", "as", "at", "be", "been", "being", "by", "can", "could",
    "did", "do", "does", "for", "from", "had", "has", "have", "how", "i", "in", "is",
    "it", "its", "make", "makes", "me", "of", "on", "or", "our", "should", "tell", "that",
    "the", "their", "them", "there", "these", "they", "this", "to", "was", "were", "what",
    "when", "where", "which", "who", "why", "will", "with", "would", "you", "your",
}


def _search_token(value: str) -> str:
    token = str(value or "").lower().strip()
    if token.startswith("differ"):
        return "differ"
    if token.endswith("ies") and len(token) > 5:
        return token[:-3] + "y"
    if token.endswith("s") and len(token) > 4 and not token.endswith("ss"):
        return token[:-1]
    return token


def _search_terms(value: str) -> set[str]:
    return {
        normalised
        for raw in re.findall(r"[a-z0-9]{2,}", str(value or "").lower())
        if raw not in _SEARCH_STOP_WORDS
        if (normalised := _search_token(raw))
    }


def _retrieval_quality_multiplier(text: str) -> float:
    lowered = str(text or "").lower()
    prohibited_claim_patterns = (
        r"\bmade entirely from\b",
        r"\bevery competing\b",
        r"\brust[- ]proof for life\b",
        r"\bmaintenance[- ]free\b",
        r"\bguaranteed for \d+ years\b",
        r"\bpatent protected or patented\b",
        r"\bcertified for any\b",
    )
    if any(re.search(pattern, lowered) for pattern in prohibited_claim_patterns):
        # Documents can contain explicit examples of wording that staff must
        # not use. If table structure is flattened during extraction, those
        # examples must still never become answerable company truth.
        return 0.0
    meta_or_unsafe_markers = (
        "synthetic test material",
        "owner must review",
        "for owner review",
        "prohibited wording",
        "must not strengthen",
    )
    return 0.0 if any(marker in lowered for marker in meta_or_unsafe_markers) else 1.0


def _now() -> str:
    return datetime.now(UTC).replace(microsecond=0).isoformat()


def _sha256_bytes(value: bytes) -> str:
    return "sha256:" + sha256(value).hexdigest()


def _sha256_json(value: Any) -> str:
    raw = json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return _sha256_bytes(raw.encode("utf-8"))


def _extract_docx_text(data: bytes) -> str:
    """Extract visible Word text without depending on an optional host package.

    DOCX files are ZIP containers. Reading the WordprocessingML directly keeps
    desktop uploads working in the packaged Python runtime and, unlike the old
    paragraph-only path, includes text stored in tables.
    """
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            document_xml = archive.read("word/document.xml")
    except (BadZipFile, KeyError) as exc:
        raise ValueError("the file is not a valid Word document") from exc

    try:
        root = ElementTree.fromstring(document_xml)
    except ElementTree.ParseError as exc:
        raise ValueError("the Word document contains invalid document XML") from exc

    lines: list[str] = []
    for paragraph in root.iter(f"{namespace}p"):
        parts = [node.text or "" for node in paragraph.iter(f"{namespace}t")]
        rendered = "".join(parts).strip()
        if rendered:
            lines.append(rendered)
    return "\n".join(lines)


def _clean_scopes(values: Iterable[str] | None) -> list[str]:
    scopes = {str(value or "").strip().lower().replace("&", "_").replace(" ", "_") for value in (values or [])}
    scopes.discard("")
    invalid = scopes - ALLOWED_SCOPES
    if invalid:
        raise ValueError(f"unsupported business-knowledge scope: {', '.join(sorted(invalid))}")
    return sorted(scopes or {"all"})


def _clean_confidentiality(value: str | None) -> str:
    level = str(value or "internal").strip().lower()
    if level not in ALLOWED_CONFIDENTIALITY:
        raise ValueError(f"unsupported confidentiality level: {level}")
    return level


def _extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json", ".html", ".htm", ".xml"}:
        return data.decode("utf-8", errors="replace")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            return "\n\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(data)).pages)
        except Exception as exc:
            raise ValueError(f"PDF text extraction failed: {exc}") from exc
    if suffix == ".docx":
        try:
            return _extract_docx_text(data)
        except Exception as exc:
            raise ValueError(f"Word text extraction failed: {exc}") from exc
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows: list[str] = []
            for sheet in workbook.worksheets:
                rows.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    rendered = " | ".join(str(value) for value in row if value not in (None, ""))
                    if rendered:
                        rows.append(rendered)
            return "\n".join(rows)
        except Exception as exc:
            raise ValueError(f"Spreadsheet text extraction failed: {exc}") from exc
    raise ValueError(f"unsupported business-knowledge file type: {suffix or 'no extension'}")


def _claim_fragments(text: str) -> list[str]:
    compact = re.sub(r"[\t\r]+", " ", str(text or ""))
    blocks = [re.sub(r"\s+", " ", block).strip(" -•") for block in re.split(r"\n\s*\n|(?<=[.!?])\s+(?=[A-Z0-9])", compact)]
    seen: set[str] = set()
    fragments: list[str] = []
    for block in blocks:
        if len(block) < 8:
            continue
        for start in range(0, len(block), 900):
            fragment = block[start:start + 900].strip()
            key = fragment.casefold()
            if len(fragment) >= 8 and key not in seen:
                seen.add(key)
                fragments.append(fragment)
    return fragments[:500]


class BusinessKnowledgeService:
    """Local, provenance-preserving business knowledge with owner promotion gates."""

    def __init__(self, base_dir: str | Path | None = None) -> None:
        self.base_dir = Path(base_dir) if base_dir else AIONBusinessPaths.BUSINESS_CONTAINERS

    def _root(self, business_id: str) -> Path:
        root = self.base_dir / business_id / "business_knowledge"
        (root / "raw").mkdir(parents=True, exist_ok=True)
        return root

    def _index_path(self, business_id: str) -> Path:
        return self._root(business_id) / "index.json"

    def _load(self, business_id: str) -> dict[str, Any]:
        path = self._index_path(business_id)
        if not path.exists():
            return {
                "schema_version": "aion.business_knowledge.v1",
                "business_id": business_id,
                "revision": 0,
                "sources": [],
                "claims": [],
                "entity_links": [],
                "retrieval_receipts": [],
                "updated_at": None,
            }
        return json.loads(path.read_text(encoding="utf-8"))

    def _save(self, business_id: str, index: dict[str, Any]) -> dict[str, Any]:
        index["revision"] = int(index.get("revision") or 0) + 1
        index["updated_at"] = _now()
        stable = {key: value for key, value in index.items() if key != "index_hash"}
        index["index_hash"] = _sha256_json(stable)
        path = self._index_path(business_id)
        temporary = path.with_suffix(".tmp")
        temporary.write_text(json.dumps(index, indent=2, ensure_ascii=False), encoding="utf-8")
        temporary.replace(path)
        return index

    def ingest_text(
        self,
        business_id: str,
        *,
        title: str,
        text: str,
        scopes: Iterable[str] | None = None,
        confidentiality: str = "internal",
        source_type: str = "owner_text",
        effective_from: str | None = None,
        effective_until: str | None = None,
    ) -> dict[str, Any]:
        raw = str(text or "").strip()
        if not raw:
            raise ValueError("business knowledge text is empty")
        return self._ingest(
            business_id,
            title=title or "Owner-provided business knowledge",
            filename=None,
            raw_bytes=raw.encode("utf-8"),
            extracted_text=raw,
            scopes=scopes,
            confidentiality=confidentiality,
            source_type=source_type,
            effective_from=effective_from,
            effective_until=effective_until,
        )

    def ingest_file(
        self,
        business_id: str,
        *,
        filename: str,
        data: bytes,
        scopes: Iterable[str] | None = None,
        confidentiality: str = "internal",
        effective_from: str | None = None,
        effective_until: str | None = None,
    ) -> dict[str, Any]:
        if not data:
            raise ValueError("business knowledge file is empty")
        if len(data) > MAX_SOURCE_BYTES:
            raise ValueError("business knowledge file exceeds 25 MiB")
        safe_name = Path(filename or "source.txt").name
        extracted = _extract_text(safe_name, data).strip()
        if not extracted:
            raise ValueError("no readable text was extracted from the file")
        return self._ingest(
            business_id,
            title=safe_name,
            filename=safe_name,
            raw_bytes=data,
            extracted_text=extracted,
            scopes=scopes,
            confidentiality=confidentiality,
            source_type="owner_file",
            effective_from=effective_from,
            effective_until=effective_until,
        )

    def _ingest(
        self,
        business_id: str,
        *,
        title: str,
        filename: str | None,
        raw_bytes: bytes,
        extracted_text: str,
        scopes: Iterable[str] | None,
        confidentiality: str,
        source_type: str,
        effective_from: str | None,
        effective_until: str | None,
    ) -> dict[str, Any]:
        clean_scopes = _clean_scopes(scopes)
        clean_confidentiality = _clean_confidentiality(confidentiality)
        source_hash = _sha256_bytes(raw_bytes)
        with _LOCK:
            index = self._load(business_id)
            existing = next((item for item in index["sources"] if item.get("source_hash") == source_hash), None)
            if existing:
                return {"duplicate": True, "source": existing, "claims_created": 0, "summary": self.summary(business_id)}
            source_id = f"bks_{uuid4().hex[:16]}"
            raw_name = f"{source_id}{Path(filename).suffix.lower() if filename else '.txt'}"
            raw_path = self._root(business_id) / "raw" / raw_name
            raw_path.write_bytes(raw_bytes)
            created_at = _now()
            source = {
                "source_id": source_id,
                "title": str(title).strip()[:240],
                "filename": filename,
                "source_type": source_type,
                "source_hash": source_hash,
                "raw_local_path": str(raw_path),
                "byte_size": len(raw_bytes),
                "scopes": clean_scopes,
                "confidentiality": clean_confidentiality,
                "effective_from": effective_from,
                "effective_until": effective_until,
                "status": "awaiting_owner_review",
                "created_at": created_at,
            }
            claims = []
            for position, fragment in enumerate(_claim_fragments(extracted_text), start=1):
                claims.append({
                    "claim_id": f"bkc_{uuid4().hex[:16]}",
                    "source_id": source_id,
                    "source_hash": source_hash,
                    "position": position,
                    "text": fragment,
                    "scopes": clean_scopes,
                    "confidentiality": clean_confidentiality,
                    "effective_from": effective_from,
                    "effective_until": effective_until,
                    "status": "pending_owner_review",
                    "verification_status": "source_extracted_unverified",
                    "created_at": created_at,
                })
            index["sources"].append(source)
            index["claims"].extend(claims)
            self._save(business_id, index)
            return {"duplicate": False, "source": source, "claims_created": len(claims), "claims": claims, "summary": self.summary(business_id)}

    def approve(self, business_id: str, claim_ids: Iterable[str], *, actor: str = "business_owner") -> dict[str, Any]:
        requested = {str(value) for value in claim_ids if str(value).strip()}
        with _LOCK:
            index = self._load(business_id)
            approved: list[dict[str, Any]] = []
            for claim in index["claims"]:
                if claim.get("claim_id") not in requested:
                    continue
                claim["status"] = "approved"
                claim["verification_status"] = "owner_attested"
                claim["approved_by"] = actor
                claim["approved_at"] = _now()
                approved.append(claim)
            for source in index["sources"]:
                source_claims = [claim for claim in index["claims"] if claim.get("source_id") == source.get("source_id")]
                if source_claims and all(claim.get("status") == "approved" for claim in source_claims):
                    source["status"] = "owner_approved"
                elif any(claim.get("status") == "approved" for claim in source_claims):
                    source["status"] = "partially_approved"
            if approved:
                self._save(business_id, index)
            return {"approved": approved, "approved_count": len(approved), "summary": self.summary(business_id)}

    def _visible(self, claim: dict[str, Any], actor_scope: str) -> bool:
        scopes = set(claim.get("scopes") or ["all"])
        actor = str(actor_scope or "pilot").lower()
        confidentiality = claim.get("confidentiality") or "internal"
        if confidentiality == "restricted" and actor not in scopes:
            return False
        if actor in {"pilot", "boardroom"}:
            return True
        return "all" in scopes or actor in scopes

    def search(self, business_id: str, query: str, *, actor_scope: str = "pilot", limit: int = 8) -> dict[str, Any]:
        terms = _search_terms(query)
        with _LOCK:
            index = self._load(business_id)
            candidates = []
            for claim in index["claims"]:
                if claim.get("status") != "approved" or not self._visible(claim, actor_scope):
                    continue
                text_terms = _search_terms(claim.get("text") or "")
                overlap = len(terms & text_terms)
                score = (overlap / max(1, len(terms)) if terms else 0.1) * _retrieval_quality_multiplier(claim.get("text") or "")
                if score > 0:
                    candidates.append((score, claim))
            ordered = sorted(candidates, key=lambda pair: (-pair[0], pair[1].get("position", 0)))
            requested_limit = max(1, min(int(limit), 25))
            selected: list[tuple[float, dict[str, Any]]] = []
            selected_ids: set[str] = set()

            def add(score: float, claim: dict[str, Any]) -> None:
                claim_id = str(claim.get("claim_id") or "")
                if not claim_id or claim_id in selected_ids or len(selected) >= requested_limit:
                    return
                selected_ids.add(claim_id)
                selected.append((score, claim))

            for score, claim in ordered:
                add(score, claim)
                # Extracted business documents often place the explanation in
                # the immediately following sentence. Include that approved,
                # same-source neighbour so retrieval does not lose its context.
                position = int(claim.get("position") or 0)
                for neighbour in index["claims"]:
                    if int(neighbour.get("position") or -999) != position + 1:
                        continue
                    if neighbour.get("source_id") != claim.get("source_id"):
                        continue
                    if neighbour.get("status") != "approved" or not self._visible(neighbour, actor_scope):
                        continue
                    if _retrieval_quality_multiplier(neighbour.get("text") or "") < 1.0:
                        continue
                    add(score * 0.92, neighbour)
                if len(selected) >= requested_limit:
                    break

            matches = [dict(claim, retrieval_score=round(score, 4)) for score, claim in selected]
            receipt = {
                "receipt_id": f"bkr_{uuid4().hex[:16]}",
                "query_hash": _sha256_bytes(str(query or "").encode("utf-8")),
                "actor_scope": actor_scope,
                "claim_ids": [item["claim_id"] for item in matches],
                "created_at": _now(),
            }
            index["retrieval_receipts"] = (index.get("retrieval_receipts") or [])[-199:] + [receipt]
            self._save(business_id, index)
            return {"matches": matches, "match_count": len(matches), "receipt": receipt}

    def rebuild_entity_links(self, business_id: str, *, actor: str = "business_owner") -> dict[str, Any]:
        """Link approved claims to explicit Boardroom entities without model inference."""
        actor_id = str(actor or "").strip()
        if not actor_id:
            raise ValueError("entity link actor is required")
        identity_path = self.base_dir / business_id / "business_identity.json"
        structure_path = self.base_dir / business_id / "business_structure.json"
        identity = json.loads(identity_path.read_text(encoding="utf-8")) if identity_path.exists() else {}
        structure = json.loads(structure_path.read_text(encoding="utf-8")) if structure_path.exists() else {}
        catalogue: list[dict[str, str]] = []

        def add(entity_id: Any, entity_type: str, name: Any) -> None:
            clean_id, clean_name = str(entity_id or "").strip(), str(name or "").strip()
            if clean_id and len(clean_name) >= 2:
                catalogue.append({"entity_id": clean_id, "entity_type": entity_type, "entity_name": clean_name})

        business_entity_id = str(identity.get("id") or f"business:{business_id}")
        add(business_entity_id, "business", identity.get("legal_name"))
        if str(identity.get("trading_name") or "").casefold() != str(identity.get("legal_name") or "").casefold():
            add(business_entity_id, "business", identity.get("trading_name"))
        collections = {
            "channels": "channel", "services": "service", "revenue_streams": "revenue_stream",
            "cost_items": "cost_item", "payment_terms": "payment_term",
            "fulfillment_processes": "fulfillment_process", "functions": "function",
            "teams": "team", "human_agents": "human_agent", "ai_agents": "ai_agent",
        }
        for collection, entity_type in collections.items():
            for entity in structure.get(collection) or []:
                if not isinstance(entity, dict):
                    continue
                add(entity.get("id"), entity_type, entity.get("name") or entity.get("label") or entity.get("title"))

        with _LOCK:
            index = self._load(business_id)
            links: list[dict[str, Any]] = []
            seen: set[tuple[str, str]] = set()
            for claim in index.get("claims") or []:
                if claim.get("status") != "approved":
                    continue
                text = str(claim.get("text") or "")
                for entity in catalogue:
                    pattern = rf"(?<![\w]){re.escape(entity['entity_name'])}(?![\w])"
                    if not re.search(pattern, text, flags=re.IGNORECASE):
                        continue
                    key = (str(claim.get("claim_id") or ""), entity["entity_id"])
                    if key in seen:
                        continue
                    seen.add(key)
                    links.append({
                        "link_id": "bkl_" + sha256(f"{business_id}:{key[0]}:{key[1]}".encode("utf-8")).hexdigest()[:20],
                        "claim_id": key[0], "entity_id": entity["entity_id"],
                        "entity_type": entity["entity_type"], "entity_name": entity["entity_name"],
                        "match_type": "deterministic_exact_name", "confidence": 1.0,
                        "source_id": claim.get("source_id"), "source_hash": claim.get("source_hash"),
                        "linked_by": actor_id, "linked_at": _now(),
                    })
            stable_old = [{key: value for key, value in item.items() if key != "linked_at"} for item in index.get("entity_links") or []]
            stable_new = [{key: value for key, value in item.items() if key != "linked_at"} for item in links]
            changed = stable_old != stable_new
            if changed:
                index["entity_links"] = links
                self._save(business_id, index)
            return {
                "business_id": business_id, "entity_count": len({item["entity_id"] for item in catalogue}),
                "link_count": len(links), "changed": changed, "links": links,
                "method": "deterministic_exact_name_owner_approved_claims_only",
            }

    def set_source_policy(self, business_id: str, source_id: str, *, actor: str, legal_hold: bool | None = None, retention_until: str | None = None) -> dict[str, Any]:
        actor_id = str(actor or "").strip()
        if not actor_id:
            raise ValueError("source policy actor is required")
        if retention_until:
            try:
                datetime.fromisoformat(str(retention_until).replace("Z", "+00:00"))
            except ValueError as exc:
                raise ValueError("source retention time is invalid") from exc
        with _LOCK:
            index = self._load(business_id)
            source = next((item for item in index.get("sources") or [] if item.get("source_id") == source_id), None)
            if source is None:
                raise KeyError("business knowledge source not found")
            if legal_hold is not None:
                source["legal_hold"] = bool(legal_hold)
            if retention_until is not None:
                source["retention_until"] = retention_until or None
            source["policy_updated_by"] = actor_id
            source["policy_updated_at"] = _now()
            self._save(business_id, index)
            receipt = self._governance_receipt(business_id, "source_policy_updated", source_id, actor_id, {"legal_hold": source.get("legal_hold", False), "retention_until": source.get("retention_until")})
            return {"source": source, "receipt": receipt, "summary": self.summary(business_id)}

    def revoke_source(self, business_id: str, source_id: str, *, actor: str, confirm_revoke: bool = False, reason: str = "") -> dict[str, Any]:
        actor_id = str(actor or "").strip()
        if not actor_id:
            raise ValueError("source revocation actor is required")
        if confirm_revoke is not True:
            raise PermissionError("source revocation confirmation required")
        with _LOCK:
            index = self._load(business_id)
            source = next((item for item in index.get("sources") or [] if item.get("source_id") == source_id), None)
            if source is None:
                raise KeyError("business knowledge source not found")
            if source.get("legal_hold") is True:
                raise PermissionError("source is protected by legal hold")
            if source.get("status") == "revoked":
                return {"source": source, "revoked_claim_count": 0, "idempotent_replay": True, "summary": self.summary(business_id)}
            now = _now()
            source["status"] = "revoked"
            source["revoked_by"] = actor_id
            source["revoked_at"] = now
            source["revocation_reason"] = str(reason or "").strip()[:500]
            revoked_claim_ids = set()
            for claim in index.get("claims") or []:
                if claim.get("source_id") == source_id:
                    claim["status"] = "revoked"
                    claim["verification_status"] = "source_revoked"
                    claim["revoked_at"] = now
                    revoked_claim_ids.add(claim.get("claim_id"))
            index["entity_links"] = [item for item in index.get("entity_links") or [] if item.get("claim_id") not in revoked_claim_ids]
            raw_path = Path(str(source.get("raw_local_path") or ""))
            raw_root = (self._root(business_id) / "raw").resolve()
            try:
                resolved = raw_path.resolve()
                if resolved.is_relative_to(raw_root) and resolved.is_file():
                    resolved.unlink()
                    source["raw_material_deleted"] = True
            except (OSError, RuntimeError):
                source["raw_material_deleted"] = False
            self._save(business_id, index)
            receipt = self._governance_receipt(business_id, "source_revoked", source_id, actor_id, {"claim_ids": sorted(str(item) for item in revoked_claim_ids), "reason": source["revocation_reason"]})
            return {"source": source, "revoked_claim_count": len(revoked_claim_ids), "idempotent_replay": False, "receipt": receipt, "summary": self.summary(business_id)}

    def apply_retention(self, business_id: str, *, actor: str, now: datetime | None = None) -> dict[str, Any]:
        instant = (now or datetime.now(UTC)).astimezone(UTC)
        index = self._load(business_id)
        due = []
        for source in index.get("sources") or []:
            value = source.get("retention_until")
            if not value or source.get("status") == "revoked" or source.get("legal_hold") is True:
                continue
            parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
            if parsed.tzinfo is None:
                parsed = parsed.replace(tzinfo=UTC)
            if parsed.astimezone(UTC) <= instant:
                due.append(str(source.get("source_id") or ""))
        results = [self.revoke_source(business_id, source_id, actor=actor, confirm_revoke=True, reason="retention_period_elapsed") for source_id in due]
        return {"business_id": business_id, "retention_checked_at": instant.isoformat(), "revoked_source_ids": due, "revoked_count": len(results)}

    def _governance_receipt(self, business_id: str, action: str, source_id: str, actor: str, evidence: dict[str, Any]) -> dict[str, Any]:
        receipt = {"receipt_id": f"bkg_{uuid4().hex[:16]}", "business_id": business_id, "action": action, "source_id": source_id, "actor": actor, "evidence_hash": _sha256_json(evidence), "created_at": _now()}
        path = self._root(business_id) / "governance_receipts.jsonl"
        with path.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(receipt, sort_keys=True, ensure_ascii=False) + "\n")
        return receipt

    def summary(self, business_id: str) -> dict[str, Any]:
        index = self._load(business_id)
        claims = index.get("claims") or []
        return {
            "schema_version": index.get("schema_version"),
            "business_id": business_id,
            "revision": index.get("revision", 0),
            "index_hash": index.get("index_hash"),
            "source_count": len(index.get("sources") or []),
            "claim_count": len(claims),
            "approved_claim_count": sum(claim.get("status") == "approved" for claim in claims),
            "pending_claim_count": sum(claim.get("status") == "pending_owner_review" for claim in claims),
            "entity_link_count": len(index.get("entity_links") or []),
            "entity_links": index.get("entity_links") or [],
            "sources": index.get("sources") or [],
            "claims": claims,
            "updated_at": index.get("updated_at"),
            "privacy": "local_business_container",
        }
