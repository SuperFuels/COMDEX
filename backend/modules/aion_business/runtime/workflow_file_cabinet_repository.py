from __future__ import annotations

import json
import base64
import csv
import io
import mimetypes
import re
import sys
from uuid import uuid4
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable

from backend.modules.aion_business.runtime.paths import AIONBusinessPaths
from backend.modules.aion_business.runtime.business_vault_runtime import (
    configured_business_vault_runtime,
    guardian_mode_requested,
)
from backend.modules.aion_business.runtime.vault_domain_store import VaultDomainStore


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


class WorkflowFileCabinetRepository:
    """
    Workspace-scoped persisted workflow file cabinet.

    This stores the desktop workflow tree under the AION Business runtime root.
    It is intentionally separate from workflow execution definitions.
    """

    @staticmethod
    def _safe_id(value: str) -> str:
        value = str(value or "").strip()
        if not value:
            return "default"
        return re.sub(r"[^a-zA-Z0-9_.-]+", "-", value).strip("-") or "default"

    @classmethod
    def cabinet_dir(cls, workspace_id: str) -> Path:
        workspace_id = cls._safe_id(workspace_id)
        AIONBusinessPaths.ensure_base_dirs()
        path = AIONBusinessPaths.ROOT / "workflow_file_cabinets" / workspace_id
        path.mkdir(parents=True, exist_ok=True)
        return path

    @classmethod
    def _guardian_store(cls) -> VaultDomainStore | None:
        if not guardian_mode_requested():
            return None
        return VaultDomainStore(
            configured_business_vault_runtime().guardian,
            actor_id="workflow-file-cabinet",
        )

    @classmethod
    def cabinet_file(cls, workspace_id: str) -> Path:
        return cls.cabinet_dir(workspace_id) / "tree.json"

    @classmethod
    def default_tree(cls, workspace_id: str) -> Dict[str, Any]:
        workspace_id = cls._safe_id(workspace_id)
        now = _now_iso()
        return {
            "version": 1,
            "workspace_id": workspace_id,
            "root": {
                "id": "root",
                "type": "folder",
                "name": "Workflows",
                "children": [],
            },
            "metadata": {
                "created_at": now,
                "updated_at": now,
                "source": "aion_business_workflow_file_cabinet",
            },
        }

    @classmethod
    def load(cls, workspace_id: str) -> Dict[str, Any]:
        workspace_id = cls._safe_id(workspace_id)
        store = cls._guardian_store()
        if store is not None:
            record = store.get_optional(
                "file_cabinet.tree",
                workspace_id=workspace_id,
                entity_id="tree",
            )
            if record is None:
                payload = cls.default_tree(workspace_id)
                store.put(
                    "file_cabinet.tree",
                    workspace_id=workspace_id,
                    entity_id="tree",
                    payload=payload,
                    expected_version=None,
                    idempotency_key=f"file-cabinet-create-{workspace_id}",
                )
                return payload
            return dict(record["payload"])
        path = cls.cabinet_file(workspace_id)

        if not path.exists():
            payload = cls.default_tree(workspace_id)
            path.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")
            return payload

        payload = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(payload, dict):
            payload = cls.default_tree(workspace_id)

        payload.setdefault("version", 1)
        payload.setdefault("workspace_id", workspace_id)
        payload.setdefault("root", {"id": "root", "type": "folder", "name": "Workflows", "children": []})
        payload.setdefault("metadata", {})
        payload["metadata"].setdefault("created_at", _now_iso())
        payload["metadata"].setdefault("updated_at", _now_iso())
        payload["metadata"].setdefault("source", "aion_business_workflow_file_cabinet")
        return payload

    @classmethod
    def save(cls, workspace_id: str, tree: Dict[str, Any]) -> Dict[str, Any]:
        workspace_id = cls._safe_id(workspace_id)

        if not isinstance(tree, dict):
            raise ValueError("Workflow file cabinet payload must be an object")

        existing_created_at = None
        store = cls._guardian_store()
        existing_record = None
        path = None if store is not None else cls.cabinet_file(workspace_id)
        if store is not None:
            existing_record = store.get_optional(
                "file_cabinet.tree", workspace_id=workspace_id, entity_id="tree"
            )
            if existing_record is not None:
                existing_created_at = ((existing_record.get("payload") or {}).get("metadata") or {}).get("created_at")
        elif path is not None and path.exists():
            try:
                existing = json.loads(path.read_text(encoding="utf-8"))
                existing_created_at = (existing.get("metadata") or {}).get("created_at")
            except Exception:
                existing_created_at = None

        payload = dict(tree)
        payload["version"] = int(payload.get("version") or 1)
        payload["workspace_id"] = workspace_id

        root = payload.get("root")

        # Desktop cabinet currently uses the live visible shape:
        # {
        #   schema_version: "aion.workflow_file_cabinet.v1",
        #   business_container: "...",
        #   folders: [...]
        # }
        #
        # Older backend shape uses:
        # {
        #   version: 1,
        #   workspace_id: "...",
        #   root: { id, type, name, children }
        # }
        #
        # Accept both. If root is absent but folders exists, wrap folders into root.children.
        if not isinstance(root, dict):
            folders = payload.get("folders")
            if isinstance(folders, list):
                root = {
                    "id": "root",
                    "type": "folder",
                    "name": "Workflows",
                    "children": folders,
                }
                payload["root"] = root
            else:
                raise ValueError("Workflow file cabinet payload requires root object or folders array")

        root.setdefault("id", "root")
        root.setdefault("type", "folder")
        root.setdefault("name", "Workflows")
        root.setdefault("children", [])

        # Preserve the desktop-native shape too, so both backend and desktop callers can use it.
        if "folders" not in payload and isinstance(root.get("children"), list):
            payload["folders"] = root["children"]
        payload.setdefault("schema_version", "aion.workflow_file_cabinet.v1")
        payload.setdefault("business_container", workspace_id)

        metadata = payload.get("metadata")
        if not isinstance(metadata, dict):
            metadata = {}

        now = _now_iso()
        metadata["created_at"] = metadata.get("created_at") or existing_created_at or now
        metadata["updated_at"] = now
        metadata["source"] = "aion_business_workflow_file_cabinet"
        payload["metadata"] = metadata

        if store is not None:
            store.put(
                "file_cabinet.tree",
                workspace_id=workspace_id,
                entity_id="tree",
                payload=payload,
                expected_version=int(existing_record["version"]) if existing_record else None,
                idempotency_key=f"file-cabinet-save-{workspace_id}-{uuid4().hex}",
            )
        else:
            assert path is not None
            path.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")
        return payload

    @classmethod
    def put_artifact(
        cls,
        workspace_id: str,
        *,
        data: bytes,
        original_name: str,
        content_type: str = "application/octet-stream",
        actor_id: str = "workflow-file-cabinet",
        object_id: str | None = None,
    ) -> Dict[str, Any]:
        """Store a File Cabinet original as an encrypted vault object.

        Callers place the returned ``vault_object_id`` in the cabinet node's
        ``target``.  Guardian mode is mandatory because plaintext app-local
        storage is not an acceptable fallback for this operation.
        """
        if not guardian_mode_requested():
            raise RuntimeError("guardian_required_for_file_cabinet_original")
        sys.audit("tessaris.vault_domain_access", "file_cabinet.originals", "put")
        workspace_id = cls._safe_id(workspace_id)
        runtime = configured_business_vault_runtime()
        resolved_object_id = object_id or f"cab-{uuid4().hex}"
        metadata = runtime.objects.put_bytes(
            object_id=resolved_object_id,
            data=bytes(data),
            original_name=original_name,
            content_type=content_type,
            idempotency_key=f"cabinet-{workspace_id}-{resolved_object_id}",
            actor_id=actor_id,
            recorded_at=_now_iso(),
        )
        return {**metadata, "vault_object_id": resolved_object_id}

    @classmethod
    def _governance_receipt(
        cls, workspace_id: str, *, operation: str, item_id: str,
        actor_id: str, details: Dict[str, Any] | None = None,
    ) -> Dict[str, Any]:
        receipt = {
            "schema_version": "aion.file_cabinet_governance.v1",
            "receipt_id": f"cabinet-receipt-{uuid4().hex}",
            "workspace_id": workspace_id,
            "operation": operation,
            "item_id": item_id,
            "actor_id": actor_id,
            "details": dict(details or {}),
            "recorded_at": _now_iso(),
        }
        store = cls._guardian_store()
        if store is None:
            raise RuntimeError("guardian_required_for_file_cabinet_governance")
        store.put(
            "file_cabinet.governance",
            workspace_id=workspace_id,
            entity_id=receipt["receipt_id"],
            payload=receipt,
            expected_version=None,
            idempotency_key=receipt["receipt_id"],
            actor_id=actor_id,
        )
        return receipt

    @classmethod
    def _children_for_parent(cls, tree: Dict[str, Any], parent_id: str) -> list[Dict[str, Any]]:
        root = tree.get("root") or {}
        root_children = tree.get("folders")
        if not isinstance(root_children, list):
            root_children = root.setdefault("children", [])
        if parent_id == "root":
            root["children"] = root_children
            tree["folders"] = root_children
            return root_children
        for node in cls._walk(root_children):
            if str(node.get("id") or "") == parent_id and node.get("type") == "folder":
                return node.setdefault("children", [])
        raise KeyError("file_cabinet_parent_not_found")

    @classmethod
    def add_artifact(
        cls, workspace_id: str, *, data: bytes, original_name: str,
        content_type: str, parent_id: str = "root", display_name: str | None = None,
        document_type: str = "original", actor_id: str = "workflow-file-cabinet",
        generated: bool = False,
    ) -> Dict[str, Any]:
        workspace_id = cls._safe_id(workspace_id)
        item_id = f"cabinet-item-{uuid4().hex}"
        stored = cls.put_artifact(
            workspace_id, data=data, original_name=original_name,
            content_type=content_type, actor_id=actor_id,
            object_id=f"cab-{uuid4().hex}",
        )
        tree = cls.load(workspace_id)
        node = {
            "id": item_id,
            "type": "artifact",
            "name": str(display_name or original_name).strip() or original_name,
            "document_type": document_type,
            "content_type": content_type,
            "byte_size": len(data),
            "status": "active",
            "source_of_truth": "business_vault",
            "deletion_policy": "append_only",
            "generated": bool(generated),
            "created_at": _now_iso(),
            "created_by": actor_id,
            "target": {"vault_object_id": stored["vault_object_id"]},
        }
        cls._children_for_parent(tree, parent_id).append(node)
        saved = cls.save(workspace_id, tree)
        receipt = cls._governance_receipt(
            workspace_id,
            operation="generated" if generated else "original_uploaded",
            item_id=item_id,
            actor_id=actor_id,
            details={"parent_id": parent_id, "vault_object_id": stored["vault_object_id"]},
        )
        return {"tree": saved, "item": node, "object": stored, "receipt": receipt}

    @classmethod
    def move_item(
        cls, workspace_id: str, item_id: str, *, parent_id: str,
        actor_id: str = "workflow-file-cabinet",
    ) -> Dict[str, Any]:
        workspace_id = cls._safe_id(workspace_id)
        tree = cls.load(workspace_id)
        roots = cls._children_for_parent(tree, "root")
        found: Dict[str, Any] | None = None
        old_parent_id: str | None = None

        def remove(nodes: list[Dict[str, Any]], owner: str) -> None:
            nonlocal found, old_parent_id
            for index, node in enumerate(list(nodes)):
                if str(node.get("id") or "") == item_id:
                    found = nodes.pop(index)
                    old_parent_id = owner
                    return
                remove(node.get("children") or [], str(node.get("id") or owner))
                if found is not None:
                    return

        remove(roots, "root")
        if found is None:
            raise KeyError("file_cabinet_item_not_found")
        if found.get("status") == "tombstoned":
            raise ValueError("file_cabinet_tombstoned_item_cannot_move")
        if parent_id == item_id or any(str(node.get("id") or "") == parent_id for node in cls._walk(found.get("children") or [])):
            raise ValueError("file_cabinet_move_cycle")
        try:
            cls._children_for_parent(tree, parent_id).append(found)
        except Exception:
            cls._children_for_parent(tree, old_parent_id or "root").append(found)
            raise
        saved = cls.save(workspace_id, tree)
        receipt = cls._governance_receipt(
            workspace_id, operation="moved", item_id=item_id, actor_id=actor_id,
            details={"from_parent_id": old_parent_id, "to_parent_id": parent_id},
        )
        return {"tree": saved, "item": found, "receipt": receipt}

    @classmethod
    def tombstone_item(
        cls, workspace_id: str, item_id: str, *, reason: str,
        actor_id: str = "workflow-file-cabinet",
    ) -> Dict[str, Any]:
        workspace_id = cls._safe_id(workspace_id)
        tree = cls.load(workspace_id)
        node = next(
            (item for item in cls._walk(cls._children_for_parent(tree, "root"))
             if str(item.get("id") or "") == item_id),
            None,
        )
        if node is None:
            raise KeyError("file_cabinet_item_not_found")
        if node.get("status") == "tombstoned":
            return {"tree": tree, "item": node, "receipt": None, "idempotent_replay": True}
        node["status"] = "tombstoned"
        node["tombstoned_at"] = _now_iso()
        node["tombstoned_by"] = actor_id
        node["tombstone_reason"] = str(reason or "").strip() or "superseded"
        saved = cls.save(workspace_id, tree)
        receipt = cls._governance_receipt(
            workspace_id, operation="tombstoned", item_id=item_id, actor_id=actor_id,
            details={"reason": node["tombstone_reason"], "vault_object_retained": True},
        )
        return {"tree": saved, "item": node, "receipt": receipt, "idempotent_replay": False}

    @staticmethod
    def _walk(nodes: Iterable[Dict[str, Any]]) -> Iterable[Dict[str, Any]]:
        for node in nodes or []:
            if not isinstance(node, dict):
                continue
            yield node
            yield from WorkflowFileCabinetRepository._walk(node.get("children") or [])

    @classmethod
    def item(cls, workspace_id: str, item_id: str) -> Dict[str, Any]:
        tree = cls.load(workspace_id)
        nodes = tree.get("folders") or (tree.get("root") or {}).get("children") or []
        for node in cls._walk(nodes):
            if str(node.get("id") or "") == str(item_id or ""):
                return node
        raise KeyError("file_cabinet_item_not_found")

    @staticmethod
    def protection(node: Dict[str, Any]) -> Dict[str, Any]:
        provenance = node.get("provenance") if isinstance(node.get("provenance"), dict) else {}
        explicit = str(node.get("deletion_policy") or "").strip().lower()
        created_by = str(provenance.get("created_by") or node.get("created_by") or "").strip().lower()
        status = str(node.get("status") or "").strip().lower()
        document_type = str(node.get("document_type") or "").strip().lower()
        source = str(node.get("source_of_truth") or "").strip().lower()

        user_removable = explicit in {"user_removable", "removable"} or (
            created_by in {"user", "user_upload", "authorised_user"}
            and status not in {"accepted", "approved", "sealed", "verified", "promoted"}
        )
        protected = not user_removable and (
            explicit in {"protected", "immutable", "append_only"}
            or node.get("type") == "business_container_artifact"
            or source in {"business_container", "boardroom", "evidence_ledger"}
            or "boardroom" in document_type
            or any(token in document_type for token in ("evidence", "receipt", "ledger", "reconciliation"))
            or status in {"accepted", "approved", "sealed", "verified", "promoted"}
        )
        return {
            "protected": bool(protected),
            "deletion_policy": "protected_record" if protected else "user_removable",
            "reason": (
                "Board, evidence or governed business record"
                if protected else "User-owned item; removable with confirmation"
            ),
        }

    @classmethod
    def _artifact_path(cls, workspace_id: str, node: Dict[str, Any]) -> Path:
        raw = str((node.get("target") or {}).get("storage_path") or "").strip()
        if not raw:
            raise FileNotFoundError("artifact_storage_path_missing")

        root = AIONBusinessPaths.ROOT.resolve()
        candidate = Path(raw)
        if candidate.is_absolute():
            resolved = candidate.resolve()
        elif raw.startswith(".runtime/AION_BUSINESS/"):
            resolved = (Path.cwd() / raw).resolve()
        else:
            resolved = (root / raw).resolve()

        try:
            resolved.relative_to(root)
        except ValueError as exc:
            raise ValueError("artifact_path_outside_business_runtime") from exc
        if not resolved.exists() or not resolved.is_file():
            raise FileNotFoundError("artifact_file_missing")
        return resolved

    @staticmethod
    def _boardroom_session_blocks(payload: Dict[str, Any]) -> list[Dict[str, Any]]:
        """Render current and legacy Board packets in one standard minutes layout."""
        blocks: list[Dict[str, Any]] = []
        minutes = payload.get("minutes_document") if isinstance(payload.get("minutes_document"), dict) else {}
        result = payload.get("boardroom_council_result") if isinstance(payload.get("boardroom_council_result"), dict) else {}
        consensus = result.get("consensus") if isinstance(result.get("consensus"), dict) else {}
        context = payload.get("business_context_packet") if isinstance(payload.get("business_context_packet"), dict) else {}

        def text(value: Any) -> str:
            return str(value or "").strip().replace("**", "").replace("__", "")

        def strings(values: Any, *, limit: int = 10) -> list[str]:
            rows = values if isinstance(values, list) else []
            output: list[str] = []
            for value in rows:
                if isinstance(value, dict):
                    label = value.get("title") or value.get("step") or value.get("member") or value.get("name") or value.get("text")
                    detail = value.get("action") or value.get("position") or value.get("view") or value.get("role") or ""
                    line = " — ".join(part for part in (text(label), text(detail)) if part)
                else:
                    line = text(value).removeprefix(">").strip()
                if line and line not in output:
                    output.append(line)
            return output[:limit]

        def heading(value: str) -> None:
            blocks.append({"type": "heading", "text": value})

        def paragraph(value: Any) -> None:
            clean = text(value)
            if clean:
                blocks.append({"type": "paragraph", "text": clean})

        def bullet_section(title: str, values: Any, *, empty: str = "None recorded.") -> None:
            heading(title)
            items = strings(values)
            if items:
                blocks.append({"type": "list", "text": "\n".join(items), "items": items})
            else:
                paragraph(empty)

        session_type = text(minutes.get("meeting_type") or payload.get("session_type") or "business_assessment").replace("_", " ").title()
        generated_at = text(minutes.get("occurred_at") or payload.get("generated_at") or "Date not recorded")
        session_id = text(minutes.get("session_id") or payload.get("session_id") or "Unknown")
        blocks.extend([
            {"type": "title", "text": text(minutes.get("title") or "Board meeting minutes")},
            {"type": "subtitle", "text": f"{session_type} · {generated_at}"},
            {"type": "meta", "items": [
                {"label": "Meeting type", "value": session_type},
                {"label": "Session record", "value": session_id},
                {"label": "Record status", "value": "Protected · governed business evidence"},
                {"label": "Decision boundary", "value": "Founder approval required"},
            ]},
        ])

        attendees = minutes.get("attendees") or payload.get("council_members") or []
        bullet_section("Attendees", attendees, empty="No Board members were recorded.")

        heading("Meeting summary")
        paragraph(minutes.get("summary") or consensus.get("consensus_summary") or consensus.get("summary") or result.get("summary") or "The Board reviewed the available governed business evidence.")

        bullet_section("Business context reviewed", minutes.get("business_context") or context.get("lines"), empty="No additional business context was recorded.")

        discussion = minutes.get("discussion") or result.get("council_responses") or []
        bullet_section("Discussion", discussion, empty="No separate Board member positions were recorded.")

        decisions = minutes.get("decisions") or consensus.get("agreements") or []
        bullet_section("Decisions", decisions, empty=text(minutes.get("decision_status") or consensus.get("decision") or "No final decision was recorded."))

        actions = minutes.get("actions") or consensus.get("recommended_plan") or consensus.get("plan") or payload.get("proposed_pilot_actions") or []
        heading("Actions")
        action_rows = []
        for item in actions if isinstance(actions, list) else []:
            if isinstance(item, dict):
                action_rows.append([
                    text(item.get("action") or item.get("title") or item.get("step") or "Action to review"),
                    text(item.get("owner") or item.get("department") or "To assign"),
                    text(item.get("status") or "Proposed · approval required"),
                ])
            elif text(item):
                action_rows.append([text(item), "To assign", "Proposed · approval required"])
        if action_rows:
            blocks.append({"type": "table", "headers": ["Action", "Owner", "Status"], "rows": action_rows[:10]})
        else:
            paragraph("No actions were recorded.")

        bullet_section("Risks", minutes.get("risks") or consensus.get("risks"), empty="No risks were recorded.")
        bullet_section("Open questions", minutes.get("open_questions") or consensus.get("open_questions") or consensus.get("disagreements"), empty="No open questions were recorded.")

        provenance = minutes.get("model_provenance") if isinstance(minutes.get("model_provenance"), dict) else {}
        providers = provenance.get("providers") if isinstance(provenance.get("providers"), list) else []
        heading("Models and record provenance")
        if providers:
            blocks.append({"type": "table", "headers": ["Provider", "Model", "Source", "Status"], "rows": [[text(item.get("name")), text(item.get("model")), text(item.get("source")), text(item.get("status"))] for item in providers if isinstance(item, dict)]})
        else:
            paragraph("Model details were not captured in this earlier meeting record.")
        paragraph(provenance.get("policy") or "Connected Board providers supplied the meeting analysis; older records may not include model-level provenance.")

        evidence = minutes.get("evidence") if isinstance(minutes.get("evidence"), dict) else {}
        blocks.append({"type": "meta", "items": [
            {"label": "Business container", "value": text(evidence.get("business_container_id") or payload.get("business_id") or "Not recorded")},
            {"label": "Business Map revision", "value": text(evidence.get("business_map_revision") or "Not recorded")},
            {"label": "Evidence hash", "value": text(evidence.get("packet_hash") or payload.get("boardroom_session_packet_hash") or "Not recorded")},
            {"label": "Minutes schema", "value": text(minutes.get("schema_version") or "Legacy record · standardised on display")},
        ]})
        blocks.append({"type": "note", "text": text(minutes.get("governance") or "Approval-gated. No live external action was performed by creating these minutes.")})
        return blocks

    @classmethod
    def preview(cls, workspace_id: str, item_id: str) -> Dict[str, Any]:
        node = cls.item(workspace_id, item_id)
        protection = cls.protection(node)
        if node.get("type") == "workflow":
            return {"ok": True, "kind": "workflow", "node": node, "protection": protection}

        vault_object_id = str((node.get("target") or {}).get("vault_object_id") or "").strip()
        if vault_object_id:
            if not guardian_mode_requested():
                raise RuntimeError("guardian_required_for_file_cabinet_original")
            sys.audit("tessaris.vault_domain_access", "file_cabinet.originals", "get")
            metadata, data = configured_business_vault_runtime().objects.get_bytes(
                object_id=vault_object_id
            )
            return cls._preview_bytes(node, protection, metadata, data)

        path = cls._artifact_path(workspace_id, node)
        suffix = path.suffix.lower()
        mime_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        result: Dict[str, Any] = {
            "ok": True,
            "kind": "unknown",
            "name": str(node.get("name") or path.name),
            "document_type": str(node.get("document_type") or suffix.lstrip(".") or "artifact"),
            "mime_type": mime_type,
            "byte_size": path.stat().st_size,
            "protection": protection,
            "node": {k: v for k, v in node.items() if k != "children"},
        }

        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            sheets = []
            for sheet in workbook.worksheets[:8]:
                rows = []
                max_row = min(max(int(sheet.max_row or 1), 1), 120)
                max_col = min(max(int(sheet.max_column or 1), 1), 30)
                for row in sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True):
                    rows.append(["" if value is None else str(value) for value in row])
                sheets.append({"name": sheet.title, "rows": rows})
            result.update(kind="spreadsheet", sheets=sheets)
        elif suffix == ".csv":
            text = path.read_text(encoding="utf-8-sig", errors="replace")
            rows = list(csv.reader(io.StringIO(text)))[:250]
            result.update(kind="spreadsheet", sheets=[{"name": path.stem, "rows": rows}])
        elif suffix == ".docx":
            from docx import Document

            doc = Document(path)
            blocks = [{"type": "paragraph", "text": p.text} for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                blocks.append({"type": "table", "rows": [[cell.text for cell in row.cells] for row in table.rows]})
            result.update(kind="document", blocks=blocks[:500])
        elif suffix == ".pdf":
            if path.stat().st_size > 20 * 1024 * 1024:
                raise ValueError("pdf_preview_too_large")
            encoded = base64.b64encode(path.read_bytes()).decode("ascii")
            result.update(kind="pdf", data_url=f"data:application/pdf;base64,{encoded}")
        elif suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg"}:
            if path.stat().st_size > 20 * 1024 * 1024:
                raise ValueError("image_preview_too_large")
            encoded = base64.b64encode(path.read_bytes()).decode("ascii")
            result.update(kind="image", data_url=f"data:{mime_type};base64,{encoded}")
        elif suffix == ".json" and str(node.get("document_type") or "").lower() == "boardroom_session":
            payload = json.loads(path.read_text(encoding="utf-8"))
            if not isinstance(payload, dict):
                raise ValueError("boardroom_session_payload_invalid")
            meeting_type = str(payload.get("session_type") or "board meeting").replace("_", " ").title()
            result.update(kind="document", name=f"Board meeting minutes — {meeting_type}", blocks=cls._boardroom_session_blocks(payload))
        elif suffix in {".json", ".jsonl", ".txt", ".md", ".log", ".yaml", ".yml", ".xml", ".html"}:
            text = path.read_text(encoding="utf-8", errors="replace")
            if suffix == ".json":
                try:
                    text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
                except json.JSONDecodeError:
                    pass
            result.update(kind="text", text=text[:500_000], truncated=len(text) > 500_000)
        else:
            result.update(kind="download", message="This file can be retained and opened with its native desktop application.")
        return result

    @classmethod
    def _preview_bytes(
        cls,
        node: Dict[str, Any],
        protection: Dict[str, Any],
        metadata: Dict[str, Any],
        data: bytes,
    ) -> Dict[str, Any]:
        name = str(node.get("name") or metadata.get("original_name") or "vault-file")
        suffix = Path(str(metadata.get("original_name") or name)).suffix.lower()
        mime_type = str(metadata.get("content_type") or mimetypes.guess_type(name)[0] or "application/octet-stream")
        result: Dict[str, Any] = {
            "ok": True,
            "kind": "unknown",
            "name": name,
            "document_type": str(node.get("document_type") or suffix.lstrip(".") or "artifact"),
            "mime_type": mime_type,
            "byte_size": len(data),
            "protection": protection,
            "node": {k: v for k, v in node.items() if k != "children"},
        }
        if suffix in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheets = []
            for sheet in workbook.worksheets[:8]:
                rows = []
                for row in sheet.iter_rows(
                    min_row=1,
                    max_row=min(max(int(sheet.max_row or 1), 1), 120),
                    max_col=min(max(int(sheet.max_column or 1), 1), 30),
                    values_only=True,
                ):
                    rows.append(["" if value is None else str(value) for value in row])
                sheets.append({"name": sheet.title, "rows": rows})
            result.update(kind="spreadsheet", sheets=sheets)
        elif suffix == ".csv":
            text = data.decode("utf-8-sig", errors="replace")
            result.update(kind="spreadsheet", sheets=[{"name": Path(name).stem, "rows": list(csv.reader(io.StringIO(text)))[:250]}])
        elif suffix == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(data))
            blocks = [{"type": "paragraph", "text": paragraph.text} for paragraph in doc.paragraphs if paragraph.text.strip()]
            for table in doc.tables:
                blocks.append({"type": "table", "rows": [[cell.text for cell in row.cells] for row in table.rows]})
            result.update(kind="document", blocks=blocks[:500])
        elif suffix == ".pdf":
            if len(data) > 20 * 1024 * 1024:
                raise ValueError("pdf_preview_too_large")
            result.update(kind="pdf", data_url="data:application/pdf;base64," + base64.b64encode(data).decode("ascii"))
        elif suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg"}:
            if len(data) > 20 * 1024 * 1024:
                raise ValueError("image_preview_too_large")
            result.update(kind="image", data_url=f"data:{mime_type};base64," + base64.b64encode(data).decode("ascii"))
        elif suffix == ".json" and str(node.get("document_type") or "").lower() == "boardroom_session":
            payload = json.loads(data.decode("utf-8"))
            if not isinstance(payload, dict):
                raise ValueError("boardroom_session_payload_invalid")
            meeting_type = str(payload.get("session_type") or "board meeting").replace("_", " ").title()
            result.update(kind="document", name=f"Board meeting minutes — {meeting_type}", blocks=cls._boardroom_session_blocks(payload))
        elif suffix in {".json", ".jsonl", ".txt", ".md", ".log", ".yaml", ".yml", ".xml", ".html"}:
            text = data.decode("utf-8", errors="replace")
            if suffix == ".json":
                try:
                    text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
                except json.JSONDecodeError:
                    pass
            result.update(kind="text", text=text[:500_000], truncated=len(text) > 500_000)
        else:
            result.update(kind="download", message="This encrypted file can be opened with its native desktop application.")
        return result
